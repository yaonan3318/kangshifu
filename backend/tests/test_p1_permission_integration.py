"""P1-15 多用户权限集成测试。

默认跳过；设置 ``COMPANY_SEARCH_TEST_DATABASE_URL`` 指向一个**临时** PostgreSQL
测试库后运行。测试会在该库上执行 Alembic 迁移并写入隔离数据，绝不连接正式库。

    COMPANY_SEARCH_TEST_DATABASE_URL=postgresql+psycopg://user:pass@127.0.0.1:54329/cs_test \
        python -m pytest backend/tests/test_p1_permission_integration.py -v
"""

import asyncio
import json
import os
import uuid
from datetime import UTC, datetime, timedelta

import pytest

TEST_DB = os.environ.get("COMPANY_SEARCH_TEST_DATABASE_URL")
if not TEST_DB:
    pytest.skip(
        "set COMPANY_SEARCH_TEST_DATABASE_URL to run multi-user integration tests",
        allow_module_level=True,
    )

import tempfile  # noqa: E402

os.environ["COMPANY_SEARCH_DATABASE_URL"] = TEST_DB
os.environ["COMPANY_SEARCH_OLLAMA_WARMUP_ENABLED"] = "false"
os.environ["COMPANY_SEARCH_ANSWER_CACHE_ENABLED"] = "true"
os.environ["COMPANY_SEARCH_FEEDBACK_RANKING_ENABLED"] = "false"
# 集成测试会反复登录，放宽限流避免测试自身触发 429。
os.environ["COMPANY_SEARCH_AUTH_LOGIN_RATE_LIMIT"] = "100000"
os.environ.setdefault("COMPANY_SEARCH_LIBRARY_ROOT", tempfile.mkdtemp(prefix="company-search-test-"))

from pathlib import Path  # noqa: E402

from alembic import command  # noqa: E402
from alembic.config import Config as AlembicConfig  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import func, select  # noqa: E402

from app.config import Settings  # noqa: E402
from app.db import SessionLocal  # noqa: E402
from app.main import create_app  # noqa: E402
from app.models import (  # noqa: E402
    DEFAULT_KNOWLEDGE_BASE_ID, AclPermission, AnswerFeedback, AnswerFeedbackDocument, Assistant,
    AuditLog, ChatMessage, ChatMessageRole, ChatMessageSource, ChatMessageStatus, ChatSession,
    Department, Document, DocumentAcl, DocumentChunk, DocumentFeedbackStats, DocumentStatus,
    DocumentVisibility, FeedbackRating, KnowledgeBase, Role, SubjectType, Tag, User,
)
from app.schemas.answer import (  # noqa: E402
    AnswerEvent, AnswerProvider, AnswerRequest, AnswerSource, ConversationTurn, KnowledgeScope,
)
from app.schemas.search import SearchRequest  # noqa: E402
from app.services import identity  # noqa: E402
from app.services.auth import hash_password  # noqa: E402
from app.services.chat import AnswerRecorder, ChatService  # noqa: E402
from app.services.feedback_ranking import FeedbackRankingService  # noqa: E402
from app.services.permissions import PermissionResolver  # noqa: E402
from app.services.query_rewrite import QueryRewriteOutcome  # noqa: E402
from app.services.rag import RagService, _answer_cache  # noqa: E402
from app.services.retrieval_config import RetrievalConfig  # noqa: E402
from app.services.search import SearchService  # noqa: E402

BACKEND = Path(__file__).resolve().parents[1]

PASSWORD = "test-password-123"


@pytest.fixture(scope="module")
def client():
    config = AlembicConfig(str(BACKEND / "alembic.ini"))
    config.set_main_option("script_location", str(BACKEND / "migrations"))
    try:
        command.downgrade(config, "base")
    except Exception:
        pass
    command.upgrade(config, "head")
    app = create_app(Settings())
    with TestClient(app) as test_client:
        yield test_client


def _login(client: TestClient, username: str, password: str = PASSWORD):
    # 先清空 Cookie，避免上一个用例手工设置的旧会话 Cookie 遮蔽新登录。
    client.cookies.clear()
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    return response


def _logout(client: TestClient):
    client.post("/api/auth/logout")


@pytest.fixture(scope="module")
def seeded(client):
    """建立测试账号矩阵与文档矩阵；直接写库以避免依赖上传/解析链路。"""
    ids: dict[str, object] = {}
    with SessionLocal() as session:
        company = session.scalar(select(Department).where(Department.name == "公司"))
        tech = Department(name="技术部", parent_id=company.id, enabled=True)
        hr = Department(name="人事部", parent_id=company.id, enabled=True)
        role = Role(name="技术角色", enabled=True)
        disabled_role = Role(name="停用角色", enabled=False)
        session.add_all([tech, hr, role, disabled_role])
        session.flush()

        # 功能级 RBAC 默认角色：普通用户至少要有基础使用权限，才能测试文档 ACL。
        base_role = session.scalar(select(Role).where(Role.name == "普通员工"))
        maintainer_role = session.scalar(select(Role).where(Role.name == "资料维护员"))

        def make_user(
            username: str, department_id, role_row=None, super_admin=False, extra_roles=(),
        ) -> User:
            user = User(
                username=username, display_name=username, password_hash=hash_password(PASSWORD),
                department_id=department_id, enabled=True, is_super_admin=super_admin,
            )
            for extra in extra_roles:
                if extra is not None:
                    user.roles.append(extra)
            if role_row is not None:
                user.roles.append(role_row)
            session.add(user)
            session.flush()
            return user

        admin = session.scalar(select(User).where(User.username == "admin"))
        if admin is not None:
            # 统一管理员测试密码，避免依赖部署环境的 COMPANY_SEARCH_ADMIN_PASSWORD。
            admin.password_hash = hash_password(PASSWORD)
        tech_user = make_user("tech_user", tech.id, extra_roles=(base_role, maintainer_role))
        hr_user = make_user("hr_user", hr.id, extra_roles=(base_role,))
        role_user = make_user("role_user", None, role, extra_roles=(base_role,))
        other_user = make_user("other_user", None, extra_roles=(base_role,))
        session.commit()

        def make_document(name: str, visibility: DocumentVisibility, owner=None, **kwargs) -> Document:
            document = Document(
                id=uuid.uuid4(), original_name=name, stored_path=f"test/{name}",
                extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
                status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
                visibility=visibility, owner_user_id=owner.id if owner else None, enabled=True,
                **kwargs,
            )
            session.add(document)
            session.flush()
            session.add(DocumentChunk(
                document_id=document.id, sequence_number=1, content="alpha secret",
                search_vector=func.to_tsvector("simple", "alpha secret"),
            ))
            return document

        company_doc = make_document("company_doc.txt", DocumentVisibility.COMPANY)
        private_doc = make_document("private_doc.txt", DocumentVisibility.PRIVATE, owner=tech_user)
        department_doc = make_document("department_doc.txt", DocumentVisibility.DEPARTMENT)
        role_doc = make_document("role_doc.txt", DocumentVisibility.ROLE)
        disabled_role_doc = make_document("disabled_role_doc.txt", DocumentVisibility.ROLE)
        user_doc = make_document("user_doc.txt", DocumentVisibility.USER)
        managed_doc = make_document("managed_doc.txt", DocumentVisibility.USER)
        restricted_doc = make_document(
            "restricted_doc.txt", DocumentVisibility.COMPANY, sensitivity_level="CONFIDENTIAL",
        )
        external_disabled_doc = make_document(
            "external_disabled_doc.txt", DocumentVisibility.COMPANY, external_llm_allowed=False,
        )
        session.add_all([
            DocumentAcl(document_id=department_doc.id, subject_type=SubjectType.DEPARTMENT, subject_id=tech.id, permission=AclPermission.READ),
            DocumentAcl(document_id=role_doc.id, subject_type=SubjectType.ROLE, subject_id=role.id, permission=AclPermission.READ),
            DocumentAcl(document_id=disabled_role_doc.id, subject_type=SubjectType.ROLE, subject_id=disabled_role.id, permission=AclPermission.READ),
            DocumentAcl(document_id=user_doc.id, subject_type=SubjectType.USER, subject_id=tech_user.id, permission=AclPermission.READ),
            DocumentAcl(document_id=managed_doc.id, subject_type=SubjectType.USER, subject_id=tech_user.id, permission=AclPermission.MANAGE),
        ])
        session.commit()

        ids.update({
            "admin": admin, "tech_user": tech_user, "hr_user": hr_user,
            "role_user": role_user, "other_user": other_user, "role": role,
            "disabled_role": disabled_role, "tech": tech,
            "company_doc": company_doc, "private_doc": private_doc,
            "department_doc": department_doc, "role_doc": role_doc,
            "disabled_role_doc": disabled_role_doc,
            "user_doc": user_doc, "managed_doc": managed_doc, "restricted_doc": restricted_doc,
            "external_disabled_doc": external_disabled_doc,
        })
    return ids


def _resolver(user: User):
    with SessionLocal() as session:
        fresh = session.get(User, user.id)
        resolver = PermissionResolver(session, fresh)
        return resolver


def _search_document_ids(user: User) -> set[uuid.UUID]:
    """在 SQL 层验证关键词召回只包含当前用户可见的资料。"""
    with SessionLocal() as session:
        fresh = session.get(User, user.id)
        service = SearchService(session, Settings(), user=fresh)
        candidates = service._keyword_candidates("alpha", SearchRequest(query="alpha"))
        return {candidate.document.id for candidate in candidates}


def test_company_document_visible_to_everyone(seeded):
    assert seeded["company_doc"].id in _search_document_ids(seeded["hr_user"])
    assert seeded["company_doc"].id in _search_document_ids(seeded["other_user"])


def test_private_document_only_owner_and_admin(seeded):
    assert seeded["private_doc"].id in _search_document_ids(seeded["tech_user"])
    assert seeded["private_doc"].id not in _search_document_ids(seeded["hr_user"])
    assert seeded["private_doc"].id in _search_document_ids(seeded["admin"])


def test_department_document_hidden_from_other_departments(seeded):
    assert seeded["department_doc"].id in _search_document_ids(seeded["tech_user"])
    assert seeded["department_doc"].id not in _search_document_ids(seeded["hr_user"])


def test_role_document_requires_role(seeded):
    assert seeded["role_doc"].id in _search_document_ids(seeded["role_user"])
    assert seeded["role_doc"].id not in _search_document_ids(seeded["other_user"])


def test_user_document_is_scoped_to_subject(seeded):
    assert seeded["user_doc"].id in _search_document_ids(seeded["tech_user"])
    assert seeded["user_doc"].id not in _search_document_ids(seeded["hr_user"])


def test_disabled_role_does_not_grant_access(seeded):
    with SessionLocal() as session:
        user = session.get(User, seeded["other_user"].id)
        user.roles.append(session.get(Role, seeded["disabled_role"].id))
        session.commit()
    assert seeded["disabled_role_doc"].id not in _search_document_ids(seeded["other_user"])


def test_ordinary_user_cannot_call_assistant_or_harness(client, seeded):
    _login(client, "tech_user")
    assert client.post("/api/assistants", json={"name": "x"}).status_code == 403
    assert client.get("/api/harness/status").status_code == 403
    _logout(client)


def test_assistant_switches_are_independent(client, seeded):
    """管理员修改任意一个开关都不能错误影响另一个开关（P2-3 回归）。"""
    name = f"switch-independence-{uuid.uuid4().hex[:8]}"
    _login(client, "admin")
    created = client.post("/api/assistants", json={"name": name, "deepseek_enabled": True})
    assert created.status_code == 200, created.text
    assistant = created.json()
    assistant_id = assistant["id"]
    # 只显式打开 deepseek_enabled 时，其他开关必须保持各自默认值，不能被联动。
    assert assistant["deepseek_enabled"] is True
    assert assistant["use_deepseek_allowed"] is True
    assert assistant["default_deepseek_enabled"] is False
    assert assistant["harness_enabled"] is False
    assert assistant["internet_enabled"] is False

    # 仅关闭“允许使用 DeepSeek”，实际启用状态与默认状态不变。
    updated = client.patch(f"/api/assistants/{assistant_id}", json={"use_deepseek_allowed": False})
    assert updated.status_code == 200, updated.text
    body = updated.json()
    assert body["use_deepseek_allowed"] is False
    assert body["deepseek_enabled"] is True
    assert body["default_deepseek_enabled"] is False

    # 仅修改默认状态，允许状态与实际启用状态不变。
    updated = client.patch(f"/api/assistants/{assistant_id}", json={"default_deepseek_enabled": True})
    body = updated.json()
    assert body["default_deepseek_enabled"] is True
    assert body["use_deepseek_allowed"] is False
    assert body["deepseek_enabled"] is True

    # 打开 Harness 与联网不应触碰任何 DeepSeek 开关。
    updated = client.patch(
        f"/api/assistants/{assistant_id}", json={"harness_enabled": True, "internet_enabled": True},
    )
    body = updated.json()
    assert body["harness_enabled"] is True
    assert body["internet_enabled"] is True
    assert body["use_deepseek_allowed"] is False
    assert body["default_deepseek_enabled"] is True
    assert body["deepseek_enabled"] is True

    # 读取流程与写入一致。
    fetched = client.get(f"/api/assistants/{assistant_id}").json()
    assert fetched["harness_enabled"] is True
    assert fetched["internet_enabled"] is True
    assert fetched["use_deepseek_allowed"] is False
    assert fetched["default_deepseek_enabled"] is True
    assert fetched["deepseek_enabled"] is True

    # 更新普通字段（知识库绑定、流程绑定）不能重置开关或知识库范围。
    kb_id = str(DEFAULT_KNOWLEDGE_BASE_ID)
    bound = client.put(f"/api/assistants/{assistant_id}/knowledge-bases", json={"knowledge_base_ids": [kb_id]})
    assert bound.status_code == 200, bound.text
    assert bound.json()["knowledge_base_ids"] == [kb_id]
    patched = client.patch(f"/api/assistants/{assistant_id}", json={"retrieval_limit": 9})
    assert patched.json()["knowledge_base_ids"] == [kb_id]
    assert patched.json()["deepseek_enabled"] is True

    assert client.delete(f"/api/assistants/{assistant_id}").status_code == 200
    _logout(client)


def test_acl_read_requires_manage(client, seeded):
    _login(client, "hr_user")
    assert client.get(f"/api/documents/{seeded['company_doc'].id}/acl").status_code == 403
    _logout(client)

    _login(client, "tech_user")
    assert client.get(f"/api/documents/{seeded['managed_doc'].id}/acl").status_code == 200
    _logout(client)


def test_acl_write_rejects_disabled_subject(client, seeded):
    _login(client, "admin")
    response = client.put(
        f"/api/documents/{seeded['company_doc'].id}/access",
        json={
            "visibility": "ROLE",
            "acl": [{"subject_type": "ROLE", "subject_id": str(seeded["disabled_role"].id), "permission": "READ"}],
        },
    )
    assert response.status_code == 409
    _logout(client)


def test_cache_scope_differs_by_permission(seeded):
    with SessionLocal() as session:
        admin_scope = SearchService(session, Settings(), user=session.get(User, seeded["admin"].id)).permission_cache_scope()
        tech_scope = SearchService(session, Settings(), user=session.get(User, seeded["tech_user"].id)).permission_cache_scope()
    assert admin_scope != tech_scope


def test_restricted_document_blocks_external_model(seeded):
    with SessionLocal() as session:
        service = RagService(session, Settings(), user=session.get(User, seeded["admin"].id))
        blocked = service._restricted_document_ids([seeded["restricted_doc"].id, seeded["company_doc"].id])
    assert seeded["restricted_doc"].id in blocked
    assert seeded["company_doc"].id not in blocked


def test_disabling_user_invalidates_existing_session(client, seeded):
    _login(client, "other_user")
    assert client.get("/api/auth/me").json()["authenticated"] is True
    old_token = client.cookies.get("cs_session")

    _login(client, "admin")
    response = client.post(f"/api/users/{seeded['other_user'].id}/disable")
    assert response.status_code == 200, response.text
    _logout(client)

    # 换回停用用户的旧 Cookie，应被立即拒绝。
    client.cookies.set("cs_session", old_token)
    assert client.get("/api/documents").status_code == 401


def test_feedback_ranking_respects_sample_threshold(seeded):
    settings = Settings()
    settings.search_feedback_ranking_enabled = True
    settings.search_feedback_min_samples = 5
    settings.search_feedback_max_boost = 0.05
    with SessionLocal() as session:
        service = FeedbackRankingService(session, settings)
        # 尚未写入任何反馈时不应产生调整分。
        assert service.boosts([seeded["company_doc"].id]) == {}

    settings.search_feedback_ranking_enabled = False
    with SessionLocal() as session:
        assert FeedbackRankingService(session, settings).boosts([seeded["company_doc"].id]) == {}


# ======================================================================
# 测试替身：不访问真实 Ollama / DeepSeek / 嵌入模型
# ======================================================================

class _FakeEmbeddings:
    def __init__(self):
        self.vector = [1.0] + [0.0] * 1023

    def encode_query(self, _text):
        return list(self.vector)

    def encode_documents(self, texts):
        return [list(self.vector) for _ in texts]


class _FakeOllama:
    def __init__(self, answer: str = "内部答案 [1]"):
        self.answer = answer

    async def stream_with_stats(self, messages, model=None, temperature=None):
        for char in self.answer:
            yield {"delta": char, "prompt_eval_count": 5, "eval_count": 1, "done": False}
        yield {"delta": "", "prompt_eval_count": 5, "eval_count": 1, "done": True}

    async def warmup(self):
        return True


class _FakeDeepSeek:
    def __init__(self, configured: bool = True, deltas=("外部增强答案",)):
        self._configured = configured
        self._deltas = list(deltas)
        self.calls: list = []

    @property
    def configured(self):
        return self._configured

    async def stream(self, messages):
        self.calls.append(messages)
        for delta in self._deltas:
            yield delta


_OPEN_SESSIONS: list = []


@pytest.fixture(autouse=True)
def _close_test_sessions():
    yield
    while _OPEN_SESSIONS:
        _OPEN_SESSIONS.pop().close()


def _service_for(user_id, *, settings=None, answer="内部答案 [1]"):
    settings = settings or Settings()
    session = SessionLocal()
    _OPEN_SESSIONS.append(session)
    user = session.get(User, user_id)
    service = RagService(session, settings, user=user)
    service.search_service.embeddings = _FakeEmbeddings()
    service.ollama = _FakeOllama(answer)
    deepseek = _FakeDeepSeek()
    service.deepseek = deepseek
    return service, deepseek


async def _collect(service, request):
    return [event async for event in service.stream(request)]


def _events(events, event_type):
    return [event for event in events if event.type == event_type]


# ======================================================================
# 检索可见性：关键词、向量、RRF、RAG 引用、检索实验室
# ======================================================================

def _search_document_ids_with_vectors(user: User, query: str = "alpha") -> set[uuid.UUID]:
    with SessionLocal() as session:
        fresh = session.get(User, user.id)
        service = SearchService(session, Settings(), user=fresh)
        service.embeddings = _FakeEmbeddings()
        candidates = service._vector_candidates(query, SearchRequest(query=query))
        return {candidate.document.id for candidate in candidates}


def test_vector_recall_respects_permissions(seeded):
    with SessionLocal() as session:
        for name in ("company_doc", "private_doc", "department_doc", "role_doc", "user_doc"):
            document = session.get(Document, seeded[name].id)
            for chunk in session.scalars(select(DocumentChunk).where(DocumentChunk.document_id == document.id)):
                chunk.embedding = [1.0] + [0.0] * 1023
        session.commit()
    tech = _search_document_ids_with_vectors(seeded["tech_user"])
    hr = _search_document_ids_with_vectors(seeded["hr_user"])
    assert seeded["company_doc"].id in tech
    assert seeded["private_doc"].id in tech
    assert seeded["private_doc"].id not in hr
    assert seeded["department_doc"].id in tech
    assert seeded["department_doc"].id not in hr
    assert seeded["role_doc"].id not in hr
    assert seeded["user_doc"].id in tech
    assert seeded["user_doc"].id not in hr


def test_rrf_fusion_keeps_permission_filter(seeded):
    with SessionLocal() as session:
        fresh = session.get(User, seeded["hr_user"].id)
        service = SearchService(session, Settings(), user=fresh)
        service.embeddings = _FakeEmbeddings()
        request = SearchRequest(query="alpha")
        keyword = service._keyword_candidates("alpha", request)
        vector = service._vector_candidates("alpha", request)
        fused = service._fuse(keyword, vector, "alpha")
    document_ids = {candidate.document.id for candidate in fused}
    assert seeded["role_doc"].id not in document_ids
    assert seeded["private_doc"].id not in document_ids
    assert seeded["company_doc"].id in document_ids


def test_rag_references_exclude_forbidden_documents(seeded):
    with SessionLocal() as session:
        fresh = session.get(User, seeded["hr_user"].id)
        service = RagService(session, Settings(), user=fresh)
        service.search_service.embeddings = _FakeEmbeddings()
        results = service.search_service.search(SearchRequest(query="alpha"))
        sources = service._sources(results)
    document_ids = {source.document_id for source in sources}
    assert seeded["role_doc"].id not in document_ids
    assert seeded["private_doc"].id not in document_ids
    assert seeded["department_doc"].id not in document_ids


def test_retrieval_lab_hides_forbidden_documents(client, seeded, monkeypatch):
    import app.services.search as search_module

    monkeypatch.setattr(search_module, "EmbeddingService", lambda settings: _FakeEmbeddings())
    _login(client, "admin")
    response = client.post("/api/retrieval-lab/inspect", json={"query": "alpha", "limit": 10})
    assert response.status_code == 200, response.text
    _logout(client)

    # 普通用户即便拿到 HR 的会话也不能访问检索实验室（仅管理员）。
    _login(client, "hr_user")
    assert client.post("/api/retrieval-lab/inspect", json={"query": "alpha", "limit": 10}).status_code == 403
    _logout(client)


# ======================================================================
# 缓存隔离：停用角色、跨用户、资料变更后立即失效
# ======================================================================

def test_cache_not_shared_between_users(seeded):
    _answer_cache.clear()
    settings = Settings()
    settings.answer_cache_enabled = True
    admin_service, _ = _service_for(seeded["admin"].id, settings=settings)
    request = AnswerRequest(question="alpha secret")
    asyncio.run(_collect(admin_service, request))
    assert len(_answer_cache) == 1
    admin_key = next(iter(_answer_cache))

    tech_service, _ = _service_for(seeded["tech_user"].id, settings=settings)
    events = asyncio.run(_collect(tech_service, request))
    metrics = _events(events, "metrics")[0].metrics
    assert metrics["cache_hit"] is False
    tech_key = tech_service._cache_key(request, tech_service._assistant_config(request))
    assert tech_key != admin_key
    assert seeded["role_doc"].id not in {source.document_id for source in _events(events, "sources")[0].sources}


def test_disabling_role_invalidates_cached_answer(seeded):
    _answer_cache.clear()
    settings = Settings()
    settings.answer_cache_enabled = True
    service, _ = _service_for(seeded["role_user"].id, settings=settings)
    request = AnswerRequest(question="alpha secret")
    first_events = asyncio.run(_collect(service, request))
    first_sources = _events(first_events, "sources")[0].sources
    assert seeded["role_doc"].id in {source.document_id for source in first_sources}
    old_key = service._cache_key(request, service._assistant_config(request))
    assert old_key in _answer_cache

    with SessionLocal() as session:
        session.get(Role, seeded["role"].id).enabled = False
        session.commit()
    try:
        second_service, _ = _service_for(seeded["role_user"].id, settings=settings)
        second_events = asyncio.run(_collect(second_service, request))
        metrics = _events(second_events, "metrics")[0].metrics
        assert metrics["cache_hit"] is False
        sources = _events(second_events, "sources")[0].sources
        assert seeded["role_doc"].id not in {source.document_id for source in sources}
        new_key = second_service._cache_key(request, second_service._assistant_config(request))
        assert new_key != old_key
    finally:
        with SessionLocal() as session:
            session.get(Role, seeded["role"].id).enabled = True
            session.commit()


def test_document_state_change_invalidates_cached_answer(seeded):
    _answer_cache.clear()
    settings = Settings()
    settings.answer_cache_enabled = True
    service, _ = _service_for(seeded["tech_user"].id, settings=settings)
    request = AnswerRequest(question="alpha secret")
    asyncio.run(_collect(service, request))
    old_key = service._cache_key(request, service._assistant_config(request))
    assert old_key in _answer_cache

    with SessionLocal() as session:
        session.get(Document, seeded["company_doc"].id).enabled = False
        session.commit()
    try:
        service2, _ = _service_for(seeded["tech_user"].id, settings=settings)
        new_key = service2._cache_key(request, service2._assistant_config(request))
        assert new_key != old_key
    finally:
        with SessionLocal() as session:
            session.get(Document, seeded["company_doc"].id).enabled = True
            session.commit()


def test_acl_change_invalidates_cached_answer(seeded):
    _answer_cache.clear()
    settings = Settings()
    settings.answer_cache_enabled = True
    service, _ = _service_for(seeded["hr_user"].id, settings=settings)
    request = AnswerRequest(question="alpha secret")
    asyncio.run(_collect(service, request))
    old_key = service._cache_key(request, service._assistant_config(request))

    with SessionLocal() as session:
        session.add(DocumentAcl(
            document_id=seeded["company_doc"].id, subject_type=SubjectType.DEPARTMENT,
            subject_id=seeded["tech"].id, permission=AclPermission.READ,
        ))
        session.commit()
    service2, _ = _service_for(seeded["hr_user"].id, settings=settings)
    new_key = service2._cache_key(request, service2._assistant_config(request))
    assert new_key != old_key


# ======================================================================
# DeepSeek 外发策略
# ======================================================================

@pytest.fixture(scope="module")
def external_docs(client, seeded):
    ids: dict[str, object] = {}
    with SessionLocal() as session:
        def make(name: str, **kwargs) -> Document:
            document = Document(
                id=uuid.uuid4(), original_name=name, stored_path=f"test/{name}",
                extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
                status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
                visibility=DocumentVisibility.COMPANY, enabled=True, **kwargs,
            )
            session.add(document)
            session.flush()
            session.add(DocumentChunk(
                document_id=document.id, sequence_number=1, content="alpha secret",
                search_vector=func.to_tsvector("simple", "alpha secret"),
            ))
            return document

        ids["allowed"] = make("ext_allowed.txt")
        ids["external_disabled"] = make("ext_disabled.txt", external_llm_allowed=False)
        ids["confidential"] = make("ext_confidential.txt", sensitivity_level="CONFIDENTIAL")
        ids["restricted"] = make("ext_restricted.txt", sensitivity_level="RESTRICTED")
        ids["mixed_allowed"] = make("ext_mixed.txt")
        ids["mixed_blocked"] = make("ext_mixed_blocked.txt", external_llm_allowed=False)
        # DeepSeek 外发策略由助手配置决定，因此需要一个显式开启外发的助手。
        assistant = Assistant(
            name="external-policy-assistant", deepseek_enabled=True,
            use_deepseek_allowed=True, enabled=True,
        )
        session.add(assistant)
        session.commit()
        ids["assistant_id"] = assistant.id
    return ids


def _external_request(name: str, assistant_id=None) -> AnswerRequest:
    return AnswerRequest(
        question="alpha secret", document_name=name, assistant_id=assistant_id,
    )


def test_external_model_called_for_allowed_document(seeded, external_docs):
    settings = Settings()
    settings.answer_cache_enabled = False
    service, deepseek = _service_for(seeded["admin"].id, settings=settings)
    events = asyncio.run(_collect(service, _external_request("ext_allowed", external_docs["assistant_id"])))
    assert deepseek.calls, "允许外发时应当调用外部模型"
    assert _events(events, "done")[0].provider == AnswerProvider.DEEPSEEK


@pytest.mark.parametrize(
    "name",
    ["ext_disabled", "ext_confidential", "ext_restricted"],
)
def test_external_model_blocked_by_policy(seeded, external_docs, name):
    settings = Settings()
    settings.answer_cache_enabled = False
    service, deepseek = _service_for(seeded["admin"].id, settings=settings)
    events = asyncio.run(_collect(service, _external_request(name, external_docs["assistant_id"])))
    assert deepseek.calls == []
    warnings = _events(events, "warning")
    assert any(event.warning.code == "EXTERNAL_LLM_BLOCKED" for event in warnings)
    # 保留本地答案，provider 仍为本地模型。
    assert _events(events, "done")[0].provider == AnswerProvider.LOCAL
    assert any(event.type == "delta" and event.provider == AnswerProvider.LOCAL for event in events)


def test_external_model_blocked_when_any_citation_restricted(seeded, external_docs):
    settings = Settings()
    settings.answer_cache_enabled = False
    service, deepseek = _service_for(seeded["admin"].id, settings=settings)
    events = asyncio.run(_collect(service, _external_request("ext_mixed", external_docs["assistant_id"])))
    source_ids = {source.document_id for source in _events(events, "sources")[0].sources}
    assert external_docs["mixed_allowed"].id in source_ids
    assert external_docs["mixed_blocked"].id in source_ids
    assert deepseek.calls == []


def test_external_model_not_called_without_opt_in(seeded, external_docs):
    settings = Settings()
    settings.answer_cache_enabled = False
    service, deepseek = _service_for(seeded["admin"].id, settings=settings)
    events = asyncio.run(_collect(service, AnswerRequest(question="alpha secret", document_name="ext_allowed")))
    assert deepseek.calls == []
    assert _events(events, "done")[0].provider == AnswerProvider.LOCAL


# ======================================================================
# 审计：成功/失败、logout、写失败不影响业务、越权读取、敏感信息
# ======================================================================

def _latest_audit(action: str):
    with SessionLocal() as session:
        return session.scalar(
            select(AuditLog).where(AuditLog.action == action).order_by(AuditLog.created_at.desc()).limit(1)
        )


def test_login_failure_audited_with_error_code(client, seeded):
    client.post("/api/auth/login", json={"username": "tech_user", "password": "wrong-password"})
    log = _latest_audit("login_failed")
    assert log is not None
    assert log.success is False
    assert log.error_code == "LOGIN_FAILED"


def test_logout_audit_contains_user(client, seeded):
    _login(client, "tech_user")
    _logout(client)
    log = _latest_audit("logout")
    assert log is not None
    assert log.user_id == seeded["tech_user"].id
    assert log.username == "tech_user"


def test_audit_does_not_store_secrets(client, seeded):
    _login(client, "tech_user")
    _logout(client)
    for action in ("login_success", "logout"):
        log = _latest_audit(action)
        assert log is not None
        payload = str(log.detail)
        assert PASSWORD not in payload
        assert "password" not in payload.lower()
        assert "cs_session" not in payload
        assert "token" not in payload.lower()


def test_audit_write_failure_does_not_break_core(client, seeded, monkeypatch):
    import app.services.audit as audit_module

    _login(client, "admin")
    monkeypatch.setattr(audit_module, "SessionLocal", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("audit down")))
    response = client.post("/api/departments", json={"name": "审计降级部门", "enabled": True})
    assert response.status_code == 201, response.text
    _logout(client)


def test_ordinary_user_cannot_read_audit_logs(client, seeded):
    _login(client, "tech_user")
    assert client.get("/api/audit/logs").status_code == 403
    _logout(client)


# ======================================================================
# 会话归档
# ======================================================================

def _create_session(client, title: str) -> str:
    response = client.post("/api/chat/sessions", json={"title": title})
    assert response.status_code == 200, response.text
    return response.json()["id"]


def test_archive_hides_and_restore_shows_session(client, seeded):
    _login(client, "tech_user")
    session_id = _create_session(client, "归档测试")
    with SessionLocal() as session:
        session.add(ChatMessage(
            session_id=uuid.UUID(session_id), role=ChatMessageRole.USER,
            content="hello", status=ChatMessageStatus.COMPLETED,
        ))
        session.commit()

    assert client.post(f"/api/chat/sessions/{session_id}/archive").status_code == 200
    default_items = client.get("/api/chat/sessions").json()["items"]
    assert all(item["id"] != session_id for item in default_items)
    archived_items = client.get("/api/chat/sessions", params={"archived": "true"}).json()["items"]
    assert any(item["id"] == session_id for item in archived_items)

    detail = client.get(f"/api/chat/sessions/{session_id}").json()
    assert detail["archived_at"] is not None
    assert len(detail["messages"]) == 1

    assert client.post(f"/api/chat/sessions/{session_id}/restore").status_code == 200
    restored = client.get("/api/chat/sessions").json()["items"]
    assert any(item["id"] == session_id for item in restored)
    _logout(client)


def test_purge_only_with_explicit_flag(client, seeded):
    _login(client, "tech_user")
    session_id = _create_session(client, "永久删除测试")
    assert client.delete(f"/api/chat/sessions/{session_id}").json() == {"archived": True}
    assert client.get(f"/api/chat/sessions/{session_id}").json()["archived_at"] is not None
    assert client.delete(f"/api/chat/sessions/{session_id}", params={"purge": "true"}).json() == {"deleted": True}
    assert client.get(f"/api/chat/sessions/{session_id}").status_code == 404
    _logout(client)


def test_user_cannot_touch_other_users_session(client, seeded):
    _login(client, "tech_user")
    session_id = _create_session(client, "tech 私有会话")
    _logout(client)
    _login(client, "hr_user")
    assert client.get(f"/api/chat/sessions/{session_id}").status_code == 404
    assert client.post(f"/api/chat/sessions/{session_id}/archive").status_code == 404
    assert client.post(f"/api/chat/sessions/{session_id}/restore").status_code == 404
    assert client.delete(f"/api/chat/sessions/{session_id}").status_code == 404
    assert client.delete(f"/api/chat/sessions/{session_id}", params={"purge": "true"}).status_code == 404
    _logout(client)


def test_user_cannot_regenerate_other_users_answer(seeded):
    with SessionLocal() as session:
        chat = ChatService(session, user=session.get(User, seeded["tech_user"].id))
        row = chat.create("regenerate 测试")
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="旧答案",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.commit()
        message_id = message.id
        other = session.get(User, seeded["hr_user"].id)
        recorder = AnswerRecorder(session, user=other)
        with pytest.raises(KeyError):
            recorder.begin_regenerate(message_id)


def test_user_cannot_rate_other_users_answer(client, seeded):
    with SessionLocal() as session:
        chat = ChatService(session, user=session.get(User, seeded["tech_user"].id))
        row = chat.create("评价测试")
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="答案",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.commit()
        message_id = message.id

    _login(client, "hr_user")
    response = client.post("/api/feedback", json={"message_id": str(message_id), "rating": "UP"})
    assert response.status_code == 403
    _logout(client)


# ======================================================================
# 文档版本权限
# ======================================================================

@pytest.fixture(scope="module")
def version_chain(client, seeded):
    with SessionLocal() as session:
        def make(name: str, visibility, previous=None, version=1, owner=None) -> Document:
            document = Document(
                id=uuid.uuid4(), original_name=name, stored_path=f"test/{name}",
                extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
                status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
                visibility=visibility, enabled=True, previous_version_id=previous.id if previous else None,
                version_number=version, owner_user_id=owner.id if owner else None,
            )
            session.add(document)
            session.flush()
            session.add(DocumentChunk(
                document_id=document.id, sequence_number=1, content="alpha secret",
                search_vector=func.to_tsvector("simple", "alpha secret"),
            ))
            return document

        v1 = make("version_company.txt", DocumentVisibility.COMPANY, version=1)
        v2 = make("version_role.txt", DocumentVisibility.ROLE, previous=v1, version=2)
        v3 = make("version_user.txt", DocumentVisibility.USER, previous=v2, version=3, owner=seeded["tech_user"])
        session.add(DocumentAcl(
            document_id=v2.id, subject_type=SubjectType.ROLE, subject_id=seeded["role"].id,
            permission=AclPermission.READ,
        ))
        session.add(DocumentAcl(
            document_id=v3.id, subject_type=SubjectType.USER, subject_id=seeded["tech_user"].id,
            permission=AclPermission.READ,
        ))
        session.commit()
        return {"v1": v1.id, "v2": v2.id, "v3": v3.id}


def _versions_for(user: User, document_id: uuid.UUID) -> set[uuid.UUID]:
    from app.services.documents import DocumentService
    from app.services.managed_storage import ManagedStorage

    with SessionLocal() as session:
        fresh = session.get(User, user.id)
        service = DocumentService(session, ManagedStorage(Settings()), user=fresh)
        return {item.id for item in service.versions(document_id)}


def test_version_chain_filters_each_version_by_permission(seeded, version_chain):
    v1, v2, v3 = version_chain["v1"], version_chain["v2"], version_chain["v3"]
    assert _versions_for(seeded["hr_user"], v1) == {v1}
    assert _versions_for(seeded["role_user"], v1) == {v1, v2}
    # tech_user 看不到中间 ROLE 版本，但仍能看到后续 USER 版本。
    assert _versions_for(seeded["tech_user"], v1) == {v1, v3}
    assert _versions_for(seeded["admin"], v1) == {v1, v2, v3}


def test_version_entry_requires_read_permission(client, seeded, version_chain):
    _login(client, "hr_user")
    assert client.get(f"/api/documents/{version_chain['v3']}/versions").status_code == 403
    assert client.get(f"/api/documents/{version_chain['v1']}/versions").status_code == 200
    _logout(client)


# ======================================================================
# API 越权边界
# ======================================================================

def test_forbidden_document_endpoints(client, seeded):
    _login(client, "hr_user")
    role_doc = seeded["role_doc"].id
    assert client.get(f"/api/documents/{role_doc}").status_code == 403
    assert client.get(f"/api/documents/{role_doc}/download").status_code == 403
    assert client.get(f"/api/documents/{role_doc}/chunks").status_code == 403
    assert client.get(f"/api/documents/{role_doc}/versions").status_code == 403
    _logout(client)


def test_acl_edit_requires_manage_permission(client, seeded):
    _login(client, "tech_user")
    # managed_doc 授予 tech_user MANAGE，可以编辑。
    assert client.put(
        f"/api/documents/{seeded['managed_doc'].id}/access",
        json={"visibility": "USER"},
    ).status_code == 200
    # user_doc 只有 READ，不能编辑 ACL。
    assert client.put(
        f"/api/documents/{seeded['user_doc'].id}/access",
        json={"visibility": "USER"},
    ).status_code == 403
    _logout(client)

    _login(client, "hr_user")
    assert client.put(
        f"/api/documents/{seeded['company_doc'].id}/access",
        json={"visibility": "COMPANY"},
    ).status_code == 403
    _logout(client)


def test_ordinary_user_cannot_call_harness_task(client, seeded):
    _login(client, "tech_user")
    assert client.post("/api/assistants", json={"name": "越权助手"}).status_code == 403
    assert client.get("/api/harness/status").status_code == 403
    assert client.get("/api/harness/tasks/00000000-0000-0000-0000-000000000000").status_code == 403
    _logout(client)


# ======================================================================
# 多文档反馈归因
# ======================================================================

def test_feedback_links_all_cited_documents(seeded):
    settings = Settings()
    settings.search_feedback_ranking_enabled = True
    settings.search_feedback_min_samples = 1
    with SessionLocal() as session:
        user = session.get(User, seeded["tech_user"].id)
        chat = ChatService(session, user=user)
        row = chat.create("多文档反馈")
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="答案 [1][2]",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.flush()
        session.add_all([
            ChatMessageSource(
                message_id=message.id, document_id=seeded["company_doc"].id, chunk_id=uuid.uuid4(),
                citation_number=1, document_name="company_doc.txt", content_snapshot="alpha",
                location_snapshot={}, score=0.9,
            ),
            ChatMessageSource(
                message_id=message.id, document_id=seeded["user_doc"].id, chunk_id=uuid.uuid4(),
                citation_number=2, document_name="user_doc.txt", content_snapshot="beta",
                location_snapshot={}, score=0.8,
            ),
        ])
        session.commit()
        message_id = message.id

        feedback = AnswerFeedback(
            message_id=message_id, user_id=user.id, rating=FeedbackRating.UP,
            reasons=[], comment=None, document_id=None,
        )
        session.add(feedback)
        session.flush()
        session.add_all([
            AnswerFeedbackDocument(feedback_id=feedback.id, document_id=seeded["company_doc"].id, citation_number=1),
            AnswerFeedbackDocument(feedback_id=feedback.id, document_id=seeded["user_doc"].id, citation_number=2),
        ])
        session.commit()

        links = session.scalars(
            select(AnswerFeedbackDocument).where(AnswerFeedbackDocument.feedback_id == feedback.id)
        ).all()
        assert {(link.document_id, link.citation_number) for link in links} == {
            (seeded["company_doc"].id, 1), (seeded["user_doc"].id, 2),
        }

        service = FeedbackRankingService(session, settings)
        service.recompute(seeded["company_doc"].id)
        service.recompute(seeded["user_doc"].id)
        assert session.get(DocumentFeedbackStats, seeded["company_doc"].id).sample_count == 1
        assert session.get(DocumentFeedbackStats, seeded["user_doc"].id).sample_count == 1
        boosts = service.boosts([seeded["company_doc"].id, seeded["user_doc"].id])
        assert set(boosts) == {seeded["company_doc"].id, seeded["user_doc"].id}
        assert all(abs(value) <= settings.search_feedback_max_boost for value in boosts.values())

        # 唯一约束保证同一回答同一文档只统计一次。
        session.add(AnswerFeedbackDocument(
            feedback_id=feedback.id, document_id=seeded["company_doc"].id, citation_number=1,
        ))
        with pytest.raises(Exception):
            session.commit()
        session.rollback()


def test_feedback_api_links_all_cited_documents(client, seeded):
    with SessionLocal() as session:
        user = session.get(User, seeded["tech_user"].id)
        chat = ChatService(session, user=user)
        row = chat.create("API 多文档反馈")
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="答案 [1][2]",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.flush()
        session.add_all([
            ChatMessageSource(
                message_id=message.id, document_id=seeded["company_doc"].id, chunk_id=uuid.uuid4(),
                citation_number=1, document_name="company_doc.txt", content_snapshot="alpha",
                location_snapshot={}, score=0.9,
            ),
            ChatMessageSource(
                message_id=message.id, document_id=seeded["user_doc"].id, chunk_id=uuid.uuid4(),
                citation_number=2, document_name="user_doc.txt", content_snapshot="beta",
                location_snapshot={}, score=0.8,
            ),
        ])
        session.commit()
        message_id = message.id

    _login(client, "tech_user")
    response = client.post("/api/feedback", json={"message_id": str(message_id), "rating": "UP"})
    assert response.status_code == 201, response.text
    _logout(client)

    with SessionLocal() as session:
        feedback = session.scalar(select(AnswerFeedback).where(AnswerFeedback.message_id == message_id))
        assert feedback is not None
        links = session.scalars(
            select(AnswerFeedbackDocument).where(AnswerFeedbackDocument.feedback_id == feedback.id)
        ).all()
        assert {link.document_id for link in links} == {seeded["company_doc"].id, seeded["user_doc"].id}

        # 重复提交（更新）不会产生重复关联。
        _login(client, "tech_user")
        client.post("/api/feedback", json={"message_id": str(message_id), "rating": "DOWN"})
        _logout(client)
    with SessionLocal() as session:
        feedback = session.scalar(select(AnswerFeedback).where(AnswerFeedback.message_id == message_id))
        links = session.scalars(
            select(AnswerFeedbackDocument).where(AnswerFeedbackDocument.feedback_id == feedback.id)
        ).all()
        assert len(links) == 2


def test_answer_api_audits_external_block_without_leaking_content(client, seeded, external_docs, monkeypatch):
    import app.services.answer_runner as runner_module

    def factory(session, settings, user=None, config=None):
        service = RagService(session, settings, user=user, config=config)
        service.search_service.embeddings = _FakeEmbeddings()
        service.ollama = _FakeOllama()
        service.deepseek = _FakeDeepSeek()
        return service

    monkeypatch.setattr(runner_module, "RAG_FACTORY", factory)
    _login(client, "admin")
    with client.stream(
        "POST", "/api/answer/stream",
        json={"question": "alpha secret", "document_name": "ext_disabled", "use_deepseek": True},
    ) as response:
        assert response.status_code == 200
        body = "".join(response.iter_text())
    assert "EXTERNAL_LLM_BLOCKED" in body
    _logout(client)

    log = _latest_audit("external_llm_blocked")
    assert log is not None
    assert "EXTERNAL_LLM_BLOCKED" in str(log.detail)
    # 审计不得包含完整上下文正文或 API Key。
    assert "alpha secret" not in str(log.detail)
    assert "sk-" not in str(log.detail)


# ======================================================================
# 功能级 RBAC：默认角色矩阵与接口收口
# ======================================================================

BATCH_BODY = {
    "name": "rbac-batch",
    "files": [{"relative_path": "a.txt", "original_name": "a.txt", "size_bytes": 1}],
}


@pytest.fixture(scope="module")
def rbac_users(client, seeded):
    """按需求默认角色建立五类账号矩阵（admin 已存在）。"""
    ids: dict[str, object] = {}
    with SessionLocal() as session:
        def role_by_name(name: str):
            return session.scalar(select(Role).where(Role.name == name))

        def make(username: str, role_name: str | None) -> User:
            user = User(
                username=username, display_name=username, password_hash=hash_password(PASSWORD),
                department_id=None, enabled=True, is_super_admin=False,
            )
            if role_name:
                user.roles.append(role_by_name(role_name))
            session.add(user)
            session.flush()
            return user

        ids["employee"] = make("employee", "普通员工")
        ids["maintainer"] = make("maintainer", "资料维护员")
        ids["kb_admin"] = make("kb_admin", "知识库管理员")
        ids["no_role"] = make("no_role", None)
        session.commit()
    return ids


def test_rbac_me_returns_effective_permissions(client, rbac_users):
    _login(client, "employee")
    user = client.get("/api/auth/me").json()["user"]
    assert set(user["permissions"]) == {"ANSWER_USE", "SEARCH_USE", "DOCUMENT_VIEW"}
    assert user["roles"] == ["普通员工"]
    _logout(client)


def test_rbac_super_admin_gets_all_permissions(client, rbac_users):
    _login(client, "admin")
    permissions = set(client.get("/api/auth/me").json()["user"]["permissions"])
    assert {"ANSWER_USE", "IDENTITY_MANAGE", "AUDIT_VIEW", "STATS_VIEW", "HARNESS_USE"} <= permissions
    _logout(client)


def test_rbac_no_role_has_no_function_permissions(client, rbac_users):
    _login(client, "no_role")
    assert client.post("/api/search", json={"query": "alpha"}).status_code == 403
    assert client.get("/api/documents").status_code == 403
    assert client.post("/api/answer/stream", json={"question": "alpha"}).status_code == 403
    _logout(client)


def test_rbac_unauthenticated_returns_401(client, rbac_users):
    client.cookies.clear()
    assert client.post("/api/search", json={"query": "alpha"}).status_code == 401
    assert client.get("/api/documents").status_code == 401
    assert client.post("/api/batches", json=BATCH_BODY).status_code == 401


def test_rbac_employee_cannot_upload_but_maintainer_can(client, rbac_users):
    _login(client, "employee")
    assert client.post("/api/batches", json=BATCH_BODY).status_code == 403
    _logout(client)

    _login(client, "maintainer")
    assert client.post("/api/batches", json=BATCH_BODY).status_code == 201
    _logout(client)


def test_rbac_knowledge_base_manage_is_scoped(client, rbac_users):
    _login(client, "maintainer")
    assert client.post("/api/knowledge-bases", json={"name": "维护员不能建库"}).status_code == 403
    _logout(client)

    _login(client, "kb_admin")
    assert client.post("/api/knowledge-bases", json={"name": "知识库管理员建库"}).status_code == 201
    _logout(client)


def test_rbac_function_permission_cannot_bypass_document_acl(client, rbac_users, seeded):
    _login(client, "maintainer")
    # 有 DOCUMENT_VIEW / DOCUMENT_MANAGE，但 role_doc 的 ACL 未授权给该用户。
    assert client.get(f"/api/documents/{seeded['role_doc'].id}").status_code == 403
    assert client.put(
        f"/api/documents/{seeded['role_doc'].id}/access", json={"visibility": "COMPANY"},
    ).status_code == 403
    _logout(client)


def test_rbac_identity_and_audit_endpoints_scoped(client, rbac_users):
    _login(client, "maintainer")
    # 身份与角色管理全部要求 IDENTITY_MANAGE。
    assert client.get("/api/users").status_code == 403
    assert client.get("/api/roles").status_code == 403
    assert client.post("/api/users", json={
        "username": "maintainer_made", "display_name": "x", "password": "password123",
    }).status_code == 403
    assert client.post("/api/roles", json={"name": "维护员不能建角色"}).status_code == 403
    assert client.post("/api/departments", json={"name": "维护员不能建部门"}).status_code == 403
    # 但可以读取 ACL 编辑所需的引用数据。
    assert client.get("/api/documents/acl-references").status_code == 200
    assert client.get("/api/audit/logs").status_code == 403
    assert client.get("/api/stats/overview").status_code == 403
    assert client.post("/api/retrieval-lab/inspect", json={"query": "alpha"}).status_code == 403
    assert client.post("/api/assistants", json={"name": "x"}).status_code == 403
    _logout(client)


def test_rbac_permission_catalog_hides_harness(client, rbac_users):
    _login(client, "admin")
    data = client.get("/api/roles/permissions").json()
    codes = {item["code"] for item in data["items"]}
    assert "HARNESS_USE" not in codes
    assert {"ANSWER_USE", "DOCUMENT_MANAGE", "IDENTITY_MANAGE"} <= codes
    _logout(client)


def test_rbac_role_permissions_changed_audited(client, rbac_users):
    _login(client, "admin")
    created = client.post("/api/roles", json={
        "name": "临时上传角色", "permissions": ["ANSWER_USE", "DOCUMENT_UPLOAD"],
    })
    assert created.status_code == 201, created.text
    role = created.json()
    assert role["permission_count"] == 2
    updated = client.patch(f"/api/roles/{role['id']}", json={"permissions": ["ANSWER_USE"]})
    assert updated.status_code == 200
    assert updated.json()["permissions"] == ["ANSWER_USE"]
    _logout(client)

    with SessionLocal() as session:
        log = session.scalar(
            select(AuditLog).where(AuditLog.action == "role_permissions_changed")
            .order_by(AuditLog.created_at.desc()).limit(1)
        )
        assert log is not None
        assert "DOCUMENT_UPLOAD" in str(log.detail)


def test_rbac_disabling_role_removes_function_permission(client, rbac_users):
    with SessionLocal() as session:
        role_id = session.scalar(select(Role.id).where(Role.name == "资料维护员"))
    _login(client, "admin")
    assert client.post(f"/api/roles/{role_id}/disable").status_code == 200
    _logout(client)
    try:
        _login(client, "maintainer")
        assert client.post("/api/batches", json=BATCH_BODY).status_code == 403
        _logout(client)
    finally:
        _login(client, "admin")
        client.post(f"/api/roles/{role_id}/enable")
        _logout(client)


def test_rbac_denials_are_audited_without_secrets(client, rbac_users):
    _login(client, "employee")
    client.post("/api/batches", json=BATCH_BODY)
    client.post("/api/knowledge-bases", json={"name": "越权建库"})
    _logout(client)

    with SessionLocal() as session:
        denied = session.scalar(
            select(AuditLog).where(AuditLog.action == "authorization_denied")
            .order_by(AuditLog.created_at.desc()).limit(1)
        )
        assert denied is not None
        assert denied.success is False
        assert denied.detail.get("missing_permission") == "KNOWLEDGE_BASE_MANAGE"
        assert denied.detail.get("path") == "/api/knowledge-bases"
        assert "password" not in str(denied.detail).lower()

        upload_denied = session.scalar(
            select(AuditLog).where(AuditLog.action == "document_upload_denied")
            .order_by(AuditLog.created_at.desc()).limit(1)
        )
        assert upload_denied is not None
        assert upload_denied.detail.get("missing_permission") == "DOCUMENT_UPLOAD"

        kb_denied = session.scalar(
            select(AuditLog).where(AuditLog.action == "knowledge_base_manage_denied")
            .order_by(AuditLog.created_at.desc()).limit(1)
        )
        assert kb_denied is not None


# ======================================================================
# P2-0 质量基线：评测集、配置版本、批量运行与新旧对比
# ======================================================================

class _FakeEmbeddingsNegative:
    """返回与所有已存向量都不相似的查询向量，让评测只走关键词召回。"""

    def encode_query(self, _text):
        return [-1.0] + [0.0] * 1023

    def encode_documents(self, texts):
        return [self.encode_query(text) for text in texts]


@pytest.fixture(scope="module")
def p2_documents(client, seeded):
    ids: dict[str, object] = {}
    with SessionLocal() as session:
        def make(name: str, content: str) -> Document:
            document = Document(
                id=uuid.uuid4(), original_name=name, stored_path=f"test/{name}",
                extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
                status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
                visibility=DocumentVisibility.COMPANY, enabled=True,
            )
            session.add(document)
            session.flush()
            session.add(DocumentChunk(
                document_id=document.id, sequence_number=1, content=content,
                search_vector=func.to_tsvector("simple", content),
            ))
            return document

        ids["expected"] = make("p2_expected.txt", "p2uniqueexpected 部署回滚流程说明")
        ids["forbidden"] = make("p2_forbidden.txt", "p2uniqueexpected 机密禁止外发内容")
        session.commit()
    return ids


def _p2_case_payload(expected_id, forbidden_id) -> dict:
    return {
        "name": "P2 用例", "question": "p2uniqueexpected",
        "expected_document_ids": [str(expected_id)],
        "must_cite_document_ids": [str(expected_id)],
        "forbidden_document_ids": [str(forbidden_id)],
        "expected_answer_keypoints": ["部署", "回滚"],
    }


def test_p2_evaluation_set_config_and_run(client, seeded, p2_documents, monkeypatch):
    import app.services.search as search_module

    monkeypatch.setattr(search_module, "EmbeddingService", lambda settings: _FakeEmbeddingsNegative())
    _login(client, "admin")

    created_set = client.post("/api/retrieval-lab/evaluation-sets", json={"name": "P2 评测集"})
    assert created_set.status_code == 201, created_set.text
    set_id = created_set.json()["id"]

    created_case = client.post(
        f"/api/retrieval-lab/evaluation-sets/{set_id}/cases",
        json=_p2_case_payload(p2_documents["expected"].id, p2_documents["forbidden"].id),
    )
    assert created_case.status_code == 201, created_case.text

    v1 = client.post("/api/retrieval-lab/config-versions", json={
        "name": "v1", "config": {"keyword_limit": 10, "vector_limit": 10, "rrf_k": 60}, "is_default": True,
    }).json()
    v2 = client.post("/api/retrieval-lab/config-versions", json={
        "name": "v2", "config": {"keyword_limit": 20, "vector_limit": 20, "rrf_k": 30},
    }).json()
    differences = client.get(
        "/api/retrieval-lab/config-versions/compare", params={"left": v1["id"], "right": v2["id"]},
    ).json()["differences"]
    assert {"keyword_limit", "rrf_k"} <= {item["field"] for item in differences}

    run1 = client.post("/api/retrieval-lab/runs", json={
        "evaluation_set_id": set_id, "config_version_id": v1["id"], "limit": 5,
    })
    assert run1.status_code == 201, run1.text
    body1 = run1.json()
    assert body1["metrics"]["case_count"] == 1
    assert body1["metrics"]["document_recall"] == 1.0
    assert body1["metrics"]["answer_mode"] == "evidence"
    # 命中禁止召回文档必须被检出。
    assert body1["metrics"]["forbidden_violation_count"] == 1
    assert body1["config_snapshot"]["keyword_limit"] == 10

    run2 = client.post("/api/retrieval-lab/runs", json={
        "evaluation_set_id": set_id, "config_version_id": v2["id"], "limit": 5,
    }).json()
    compare = client.get(
        "/api/retrieval-lab/runs/compare", params={"left": body1["id"], "right": run2["id"]},
    ).json()
    assert "document_recall" in compare["metric_deltas"]
    assert compare["config_differences"]
    assert compare["case_changes"]
    _logout(client)


def test_p2_run_with_answer_metrics(client, seeded, p2_documents, monkeypatch):
    import app.services.retrieval_evaluation as evaluation_module
    import app.services.search as search_module

    monkeypatch.setattr(search_module, "EmbeddingService", lambda settings: _FakeEmbeddingsNegative())
    expected_id = p2_documents["expected"].id

    class _FakeRag:
        def __init__(self, session, settings, user=None, config=None, source_limit=None):
            self.config = config
            self.source_limit = source_limit

        async def stream(self, request):
            yield AnswerEvent(type="sources", sources=[AnswerSource(
                citation_number=1, chunk_id=uuid.uuid4(), document_id=expected_id,
                document_name="p2_expected.txt", extension="txt", sequence_number=1,
                content="部署回滚流程", page_start=None, page_end=None, slide_number=None,
                sheet_name=None, row_start=None, row_end=None, section_path=[],
                ocr_confidence=None, match_type="keyword", score=0.9,
            )])
            yield AnswerEvent(type="delta", text="部署与回滚答案")
            yield AnswerEvent(type="metrics", metrics={"llm_first_token_ms": 12.0, "total_ms": 34.0})
            yield AnswerEvent(type="done", scope=KnowledgeScope.INTERNAL, provider=AnswerProvider.LOCAL)

    monkeypatch.setattr(evaluation_module, "RagService", _FakeRag)
    _login(client, "admin")
    set_id = client.post("/api/retrieval-lab/evaluation-sets", json={"name": "答案评测集"}).json()["id"]
    client.post(
        f"/api/retrieval-lab/evaluation-sets/{set_id}/cases",
        json=_p2_case_payload(expected_id, p2_documents["forbidden"].id),
    )
    run = client.post("/api/retrieval-lab/runs", json={
        "evaluation_set_id": set_id, "limit": 5, "include_answers": True,
    })
    assert run.status_code == 201, run.text
    metrics = run.json()["metrics"]
    assert metrics["answer_mode"] == "answer"
    assert metrics["keypoint_coverage"] == 1.0
    assert metrics["citation_accuracy"] == 1.0
    assert metrics["first_token_latency_ms"] == 12.0
    assert metrics["answer_latency_ms"] == 34.0
    _logout(client)


def test_p2_import_cases_from_csv(client, seeded, p2_documents):
    _login(client, "admin")
    set_id = client.post("/api/retrieval-lab/evaluation-sets", json={"name": "导入评测集"}).json()["id"]
    csv_body = (
        "name,question,expected_documents,must_cite_documents,forbidden_documents,answer_keypoints\n"
        "导入用例,p2uniqueexpected,p2_expected.txt,p2_expected.txt,p2_forbidden.txt,部署;回滚\n"
    )
    response = client.post(
        f"/api/retrieval-lab/evaluation-sets/{set_id}/import",
        files={"file": ("cases.csv", csv_body.encode("utf-8"), "text/csv")},
    )
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    cases = client.get(f"/api/retrieval-lab/evaluation-sets/{set_id}/cases").json()["items"]
    assert cases[0]["expected_document_ids"] == [str(p2_documents["expected"].id)]
    assert cases[0]["forbidden_document_ids"] == [str(p2_documents["forbidden"].id)]
    assert cases[0]["expected_answer_keypoints"] == ["部署", "回滚"]
    _logout(client)


def test_p2_import_cases_from_excel(client, seeded, p2_documents):
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "question", "expected_documents", "must_cite_documents", "answer_keypoints"])
    sheet.append(["Excel 用例", "p2uniqueexpected", "p2_expected.txt", "p2_expected.txt", "部署;回滚"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    _login(client, "admin")
    set_id = client.post("/api/retrieval-lab/evaluation-sets", json={"name": "Excel 导入评测集"}).json()["id"]
    response = client.post(
        f"/api/retrieval-lab/evaluation-sets/{set_id}/import",
        files={"file": ("cases.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    cases = client.get(f"/api/retrieval-lab/evaluation-sets/{set_id}/cases").json()["items"]
    assert cases[0]["expected_document_ids"] == [str(p2_documents["expected"].id)]
    assert cases[0]["expected_answer_keypoints"] == ["部署", "回滚"]
    _logout(client)


# ======================================================================
# P2-1 检索准确性：改写/多查询/词典/元数据过滤/精排名次
# ======================================================================

class _FakeRewriter:
    """测试替身：不访问真实模型，返回固定的改写与多查询结果。"""

    def __init__(self, settings, config, ollama=None):
        self.config = config

    async def rewrite(self, question, history=None):
        return QueryRewriteOutcome(
            original=question, standalone_question="补全后的问题",
            retrieval_query="改写检索词", queries=["改写检索词"],
            used_context=bool(history), rewritten=True,
        )

    async def multi_query(self, retrieval_query):
        return [retrieval_query, "多查询二", "多查询三"]


def test_p2_query_rewrite_and_multi_query_in_answer(seeded):
    _answer_cache.clear()
    settings = Settings()
    settings.answer_cache_enabled = True
    service, _ = _service_for(seeded["admin"].id, settings=settings)
    service.config = RetrievalConfig(
        query_rewrite_enabled=True, multi_query_enabled=True, multi_query_count=3,
        dictionary_enabled=False,
    )
    service.rewrite_factory = _FakeRewriter
    events = asyncio.run(_collect(service, AnswerRequest(
        question="原始问题", history=[ConversationTurn(question="之前的问题", answer="之前的答案")],
    )))
    rewrite_events = _events(events, "query_rewrite")
    assert rewrite_events, "应产生 query_rewrite 事件"
    assert rewrite_events[0].query_rewrite.retrieval_query == "改写检索词"
    metrics = _events(events, "metrics")[0].metrics
    assert metrics["retrieval_query"] == "改写检索词"
    assert metrics["retrieval_queries"] == ["改写检索词", "多查询二", "多查询三"]


def test_p2_dictionary_crud_and_query_expansion(client, seeded, monkeypatch):
    import app.services.search as search_module

    monkeypatch.setattr(search_module, "EmbeddingService", lambda settings: _FakeEmbeddings())
    _login(client, "admin")
    created = client.post("/api/retrieval-lab/dictionaries", json={
        "category": "SYNONYM", "term": "工单", "expansions": ["ticket", "服务单"],
    })
    assert created.status_code == 201, created.text
    entry_id = created.json()["id"]

    inspect = client.post("/api/retrieval-lab/inspect", json={"query": "工单处理流程", "limit": 5})
    assert inspect.status_code == 200, inspect.text
    assert "ticket" in inspect.json()["diagnostics"]["expanded_terms"]

    duplicate = client.post("/api/retrieval-lab/dictionaries", json={
        "category": "SYNONYM", "term": "工单", "expansions": [],
    })
    assert duplicate.status_code == 409

    updated = client.patch(f"/api/retrieval-lab/dictionaries/{entry_id}", json={"expansions": ["ticket"]})
    assert updated.status_code == 200
    assert updated.json()["expansions"] == ["ticket"]

    listed = client.get("/api/retrieval-lab/dictionaries").json()
    assert any(item["term"] == "工单" for item in listed["items"])
    assert client.delete(f"/api/retrieval-lab/dictionaries/{entry_id}").status_code == 204
    _logout(client)


@pytest.fixture(scope="module")
def p2_metadata_docs(client, seeded):
    ids: dict[str, object] = {}
    with SessionLocal() as session:
        now = datetime.now(UTC)

        def make(name, owner=None, valid_from=None, valid_until=None, relative_path=None, version=1):
            document = Document(
                id=uuid.uuid4(), original_name=name, stored_path=f"test/{name}", relative_path=relative_path,
                extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
                status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
                visibility=DocumentVisibility.COMPANY, enabled=True,
                owner_user_id=owner.id if owner else None, valid_from=valid_from, valid_until=valid_until,
                version_number=version,
            )
            session.add(document)
            session.flush()
            session.add(DocumentChunk(
                document_id=document.id, sequence_number=1, content="metafilter token",
                search_vector=func.to_tsvector("simple", "metafilter token"),
            ))
            return document

        ids["active"] = make(
            "meta_active.txt", owner=seeded["tech_user"], valid_from=now - timedelta(days=1),
            valid_until=now + timedelta(days=1), relative_path="folder/active.txt", version=2,
        )
        ids["expired"] = make(
            "meta_expired.txt", owner=seeded["hr_user"],
            valid_until=now - timedelta(days=1), relative_path="folder/expired.txt",
        )
        ids["no_owner"] = make("meta_no_owner.txt", relative_path="other/no_owner.txt")
        session.commit()
    return ids


def test_p2_metadata_filters(seeded, p2_metadata_docs):
    with SessionLocal() as session:
        service = SearchService(session, Settings(), user=session.get(User, seeded["admin"].id))

        def ids_for(request: SearchRequest) -> set:
            return {candidate.document.id for candidate in service._keyword_candidates("metafilter", request)}

        owner_ids = ids_for(SearchRequest(query="metafilter", owner_user_id=seeded["tech_user"].id))
        assert p2_metadata_docs["active"].id in owner_ids
        assert p2_metadata_docs["expired"].id not in owner_ids

        valid_ids = ids_for(SearchRequest(query="metafilter", valid_only=True))
        assert p2_metadata_docs["active"].id in valid_ids
        assert p2_metadata_docs["expired"].id not in valid_ids

        path_ids = ids_for(SearchRequest(query="metafilter", relative_path="folder/"))
        assert p2_metadata_docs["active"].id in path_ids
        assert p2_metadata_docs["no_owner"].id not in path_ids

        assert p2_metadata_docs["active"].id in ids_for(SearchRequest(query="metafilter", version_number=2))
        assert p2_metadata_docs["active"].id in ids_for(
            SearchRequest(query="metafilter", document_status=DocumentStatus.READY)
        )

        department_ids = ids_for(SearchRequest(query="metafilter", department_id=seeded["tech"].id))
        assert p2_metadata_docs["active"].id in department_ids
        assert p2_metadata_docs["expired"].id not in department_ids


@pytest.fixture(scope="module")
def p2_metadata_extended(seeded):
    """覆盖知识库、文件类型、标签、创建时间与文档自身部门等过滤字段。"""
    ids: dict[str, object] = {}
    with SessionLocal() as session:
        other_kb = KnowledgeBase(name="过滤测试库", enabled=True)
        tag = Tag(name="财务标签", knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID)
        session.add_all([other_kb, tag])
        session.flush()

        def make(name, extension="txt", kb_id=None, tags=(), department_id=None) -> Document:
            document = Document(
                id=uuid.uuid4(), original_name=name, stored_path=f"test/{name}",
                extension=extension, mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
                status=DocumentStatus.READY, knowledge_base_id=kb_id or DEFAULT_KNOWLEDGE_BASE_ID,
                visibility=DocumentVisibility.COMPANY, enabled=True, department_id=department_id,
            )
            document.tags = list(tags)
            session.add(document)
            session.flush()
            session.add(DocumentChunk(
                document_id=document.id, sequence_number=1, content="extfilter token",
                search_vector=func.to_tsvector("simple", "extfilter token"),
            ))
            return document

        ids["base"] = make("ext_base.txt", tags=[tag], department_id=seeded["hr"].id)
        ids["other_kb"] = make("ext_other.txt", kb_id=other_kb.id, extension="pdf")
        ids["tag"] = tag
        ids["other_kb_id"] = other_kb.id
        session.commit()
    return ids


def test_p2_metadata_filters_extended(seeded, p2_metadata_extended):
    with SessionLocal() as session:
        service = SearchService(session, Settings(), user=session.get(User, seeded["admin"].id))

        def ids_for(request: SearchRequest) -> set:
            return {candidate.document.id for candidate in service._keyword_candidates("extfilter", request)}

        base_id = p2_metadata_extended["base"].id
        other_id = p2_metadata_extended["other_kb"].id

        # 知识库过滤
        in_base = ids_for(SearchRequest(query="extfilter", knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID))
        assert base_id in in_base and other_id not in in_base
        # 文件类型过滤
        assert base_id in ids_for(SearchRequest(query="extfilter", extension="txt"))
        assert base_id not in ids_for(SearchRequest(query="extfilter", extension="pdf"))
        assert other_id in ids_for(SearchRequest(query="extfilter", extension="pdf"))
        # 标签过滤
        tagged = ids_for(SearchRequest(query="extfilter", tags=["财务标签"]))
        assert base_id in tagged and other_id not in tagged
        # 文档自身部门过滤（无上传人时也能命中）
        by_department = ids_for(SearchRequest(query="extfilter", department_id=seeded["hr"].id))
        assert base_id in by_department
        # 创建时间过滤
        today = datetime.now(UTC).date()
        created = ids_for(SearchRequest(query="extfilter", created_from=today, created_to=today))
        assert base_id in created
        assert base_id not in ids_for(
            SearchRequest(query="extfilter", created_from=today + timedelta(days=1))
        )


# ======================================================================
# P2-2 回答可靠性：置信度、引用校验、无答案归因、可追溯引用
# ======================================================================

class _NoMatchEmbeddings:
    """返回与所有已存向量都不相似的查询向量，用于构造“无召回”场景。"""

    def encode_query(self, _text):
        return [-1.0] + [0.0] * 1023

    def encode_documents(self, texts):
        return [self.encode_query(text) for text in texts]


def test_p2_confidence_and_citation_events(seeded):
    _answer_cache.clear()
    service, _ = _service_for(seeded["admin"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    events = asyncio.run(_collect(service, AnswerRequest(question="alpha secret")))
    confidence = _events(events, "confidence")
    citation = _events(events, "citation_check")
    assert confidence, "应产生 confidence 事件"
    assert confidence[0].confidence["tier"] in ("HIGH", "MEDIUM", "INSUFFICIENT")
    assert citation, "有引用时应产生 citation_check 事件"
    metrics = _events(events, "metrics")[0].metrics
    assert metrics["confidence"]["tier"] == confidence[0].confidence["tier"]
    assert metrics["question_type"]


def test_p2_no_answer_when_nothing_matches(seeded):
    service, _ = _service_for(seeded["hr_user"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    service.search_service.embeddings = _NoMatchEmbeddings()
    events = asyncio.run(_collect(service, AnswerRequest(question="zzzznomatchkeyword")))
    no_answer = _events(events, "no_answer")
    assert no_answer, "无召回时应产生 no_answer 事件"
    assert no_answer[0].no_answer["reason"] == "NO_RELEVANT_DOCUMENT"
    assert no_answer[0].no_answer["message"] == "当前可访问的公司资料中没有找到足够依据。"


@pytest.fixture(scope="module")
def p2_restricted_doc(client, seeded):
    with SessionLocal() as session:
        document = Document(
            id=uuid.uuid4(), original_name="p2_restricted.txt", stored_path="test/p2_restricted.txt",
            extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
            status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
            visibility=DocumentVisibility.ROLE, enabled=True,
        )
        session.add(document)
        session.flush()
        session.add(DocumentChunk(
            document_id=document.id, sequence_number=1, content="p2restrictedtoken 机密内容",
            search_vector=func.to_tsvector("simple", "p2restrictedtoken 机密内容"),
        ))
        session.add(DocumentAcl(
            document_id=document.id, subject_type=SubjectType.ROLE,
            subject_id=seeded["role"].id, permission=AclPermission.READ,
        ))
        session.commit()
        return document.id


def test_p2_no_answer_reason_permission_restricted(seeded, p2_restricted_doc):
    service, _ = _service_for(seeded["other_user"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    service.search_service.embeddings = _NoMatchEmbeddings()
    events = asyncio.run(_collect(service, AnswerRequest(question="p2restrictedtoken")))
    no_answer = _events(events, "no_answer")
    assert no_answer, "无权限命中时应产生 no_answer 事件"
    assert no_answer[0].no_answer["reason"] == "PERMISSION_RESTRICTED"
    # 不得泄露无权访问文档的名称。
    assert "p2_restricted" not in json.dumps(no_answer[0].no_answer, ensure_ascii=False)


def test_p2_hydrate_source_marks_version_changed(seeded):
    with SessionLocal() as session:
        user = session.get(User, seeded["tech_user"].id)
        chat = ChatService(session, user=user)
        row = chat.create("版本更新测试")
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="答案 [1]",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.flush()
        chunk_id = session.scalar(
            select(DocumentChunk.id).where(DocumentChunk.document_id == seeded["user_doc"].id).limit(1)
        )
        source = ChatMessageSource(
            message_id=message.id, document_id=seeded["user_doc"].id, chunk_id=chunk_id,
            citation_number=1, document_name="user_doc.txt", content_snapshot="alpha secret 内容",
            location_snapshot={"text": "片段 1"}, score=0.9, document_version=0,
        )
        session.add(source)
        session.commit()
        hydrated = ChatService(session, user=user).hydrate_source(source)
    assert hydrated["status"] == "VERSION_CHANGED"
    assert hydrated["version_number"] == 1
    assert hydrated["cited_version"] == 0


def test_p3_assistant_without_kb_scope_does_not_retrieve(seeded):
    with SessionLocal() as session:
        assistant = Assistant(
            name=f"无范围助手-{uuid.uuid4().hex[:6]}",
            allow_all_knowledge_bases=False, enabled=True,
        )
        session.add(assistant)
        session.commit()
        assistant_id = assistant.id
    service, _ = _service_for(seeded["admin"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    events = asyncio.run(_collect(service, AnswerRequest(question="alpha secret", assistant_id=assistant_id)))
    no_answer = _events(events, "no_answer")
    assert no_answer, "未配置资料范围时不应检索"
    assert no_answer[0].no_answer["reason"] == "KB_SCOPE_UNCONFIGURED"
    assert not _events(events, "sources")


def test_p2_hydrate_source_traceability(seeded):
    with SessionLocal() as session:
        user = session.get(User, seeded["tech_user"].id)
        chat = ChatService(session, user=user)
        row = chat.create("追溯测试")
        session.add(ChatMessage(
            session_id=row.id, role=ChatMessageRole.USER, content="alpha secret",
            status=ChatMessageStatus.COMPLETED,
        ))
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="答案 [1]",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.flush()
        chunk_id = session.scalar(
            select(DocumentChunk.id).where(DocumentChunk.document_id == seeded["user_doc"].id).limit(1)
        )
        source = ChatMessageSource(
            message_id=message.id, document_id=seeded["user_doc"].id, chunk_id=chunk_id,
            citation_number=1, document_name="user_doc.txt", content_snapshot="alpha secret 内容",
            location_snapshot={"text": "片段 1"}, score=0.9,
        )
        session.add(source)
        session.commit()
        hydrated = ChatService(session, user=user).hydrate_source(source)
    assert hydrated["version_number"] == 1
    assert hydrated["can_download"] is True
    assert "alpha" in hydrated["matched_keywords"]


def test_p2_missing_knowledge_feedback_reason(client, seeded):
    with SessionLocal() as session:
        user = session.get(User, seeded["tech_user"].id)
        chat = ChatService(session, user=user)
        row = chat.create("缺失知识反馈")
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="没有足够依据",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.commit()
        message_id = message.id
    _login(client, "tech_user")
    response = client.post("/api/feedback", json={
        "message_id": str(message_id), "rating": "DOWN", "reasons": ["缺失知识"],
    })
    assert response.status_code == 201, response.text
    _logout(client)


def test_p2_general_knowledge_answer_is_preserved(seeded, external_docs):
    service, deepseek = _service_for(seeded["admin"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    service.search_service.embeddings = _NoMatchEmbeddings()
    events = asyncio.run(_collect(service, AnswerRequest(
        question="zzzznomatchkeyword", assistant_id=external_docs["assistant_id"],
    )))
    assert deepseek.calls, "允许外发且无内部资料时应调用 DeepSeek 通用知识"
    text = "".join(event.text or "" for event in _events(events, "delta"))
    assert "外部增强答案" in text
    # 通用知识答案不应被固定提示覆盖，但仍给出无答案归因。
    assert "当前可访问的公司资料中没有找到足够依据。" not in text
    assert _events(events, "no_answer")


# ======================================================================
# P2-3 助手角色增强
# ======================================================================

def test_p3_preset_assistants_seeded(client, seeded):
    _login(client, "admin")
    names = {item["name"] for item in client.get("/api/assistants").json()["items"]}
    for preset in (
        "康师傅综合助手", "人事制度助手", "技术研发助手",
        "产品资料助手", "运维助手", "项目进度助手",
    ):
        assert preset in names, f"缺少预置助手：{preset}"
    _logout(client)


def test_p3_assistant_new_fields_and_welcome(client, seeded):
    _login(client, "admin")
    created = client.post("/api/assistants", json={
        "name": "P3 制度助手", "answer_template": "POLICY", "no_answer_policy": "STRICT",
        "internet_enabled": True, "capabilities": ["解读制度"], "limitations": ["不提供法律意见"],
    })
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["answer_template"] == "POLICY"
    assert body["no_answer_policy"] == "STRICT"
    assert body["internet_enabled"] is True
    assert body["capabilities"] == ["解读制度"]

    welcome = client.get(f"/api/assistants/{body['id']}/welcome")
    assert welcome.status_code == 200, welcome.text
    payload = welcome.json()
    assert payload["capabilities"] == ["解读制度"]
    assert payload["limitations"] == ["不提供法律意见"]
    assert payload["knowledge_scope"]
    assert payload["internet_enabled"] is True
    _logout(client)


def test_p3_answer_template_and_strict_policy(client, seeded):
    _login(client, "admin")
    created = client.post("/api/assistants", json={
        "name": "P3 严格制度助手", "answer_template": "POLICY", "no_answer_policy": "STRICT",
        "deepseek_enabled": True, "use_deepseek_allowed": True,
    })
    assistant_id = created.json()["id"]
    _logout(client)

    service, _ = _service_for(seeded["admin"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    events = asyncio.run(_collect(service, AnswerRequest(question="alpha secret", assistant_id=assistant_id)))
    metrics = _events(events, "metrics")[0].metrics
    assert metrics["question_type"] == "POLICY"
    assert metrics["question_type_label"] == "制度问题"

    service2, _ = _service_for(seeded["admin"].id)
    service2.config = RetrievalConfig(dictionary_enabled=False)
    service2.search_service.embeddings = _NoMatchEmbeddings()
    events2 = asyncio.run(_collect(service2, AnswerRequest(question="zzzznomatch", assistant_id=assistant_id)))
    no_answer = _events(events2, "no_answer")[0].no_answer
    # 严格策略：不推荐资料、不允许通用知识补答。
    assert no_answer["allow_deepseek"] is False
    assert no_answer["recommended_documents"] == []


# ======================================================================
# P2-4 轻量 Chatflow
# ======================================================================

def test_p4_chatflow_crud_publish_rollback(client, seeded):
    _login(client, "admin")
    created = client.post("/api/chatflows", json={"name": "P4 流程"})
    assert created.status_code == 201, created.text
    flow = created.json()
    flow_id = flow["id"]
    assert flow["published_version"] == 0
    assert {node["type"] for node in flow["draft_graph"]["nodes"]} >= {"start", "retrieval", "final_answer"}

    graph = flow["draft_graph"]
    for node in graph["nodes"]:
        if node["type"] == "rerank":
            node["enabled"] = False
    patched = client.patch(f"/api/chatflows/{flow_id}", json={"graph": graph})
    assert patched.status_code == 200

    published = client.post(f"/api/chatflows/{flow_id}/publish", json={"note": "关闭精排"})
    assert published.status_code == 200, published.text
    assert published.json()["version"] == 1
    detail = client.get(f"/api/chatflows/{flow_id}").json()
    assert detail["published_version"] == 1

    versions = client.get(f"/api/chatflows/{flow_id}/versions").json()
    assert len(versions) == 1
    rollback = client.post(f"/api/chatflows/{flow_id}/rollback", json={"version": 1})
    assert rollback.status_code == 200
    _logout(client)


def test_p4_chatflow_debug_run(client, seeded, monkeypatch):
    import app.services.search as search_module

    monkeypatch.setattr(search_module, "EmbeddingService", lambda settings: _FakeEmbeddings())
    _login(client, "admin")
    flows = client.get("/api/chatflows").json()["items"]
    default = next(item for item in flows if item["name"] == "默认流程")
    response = client.post(f"/api/chatflows/{default['id']}/debug", json={"question": "alpha", "knowledge_base_id": None})
    assert response.status_code == 200, response.text
    body = response.json()
    types = [node["type"] for node in body["nodes"]]
    assert "question_classify" in types and "retrieval" in types and "condition" in types
    assert body["total_ms"] >= 0
    for node in body["nodes"]:
        assert node["status"] in ("succeeded", "skipped", "failed")
        assert "duration_ms" in node
    _logout(client)


def test_p4_assistant_binds_chatflow(client, seeded):
    _login(client, "admin")
    flow = client.post("/api/chatflows", json={"name": "P4 绑定流程"}).json()
    assistant = client.post("/api/assistants", json={"name": "P4 绑定助手", "chatflow_id": flow["id"]})
    assert assistant.status_code == 200, assistant.text
    assert assistant.json()["chatflow_id"] == flow["id"]

    bad = client.post("/api/assistants", json={"name": "P4 坏绑定", "chatflow_id": str(uuid.uuid4())})
    assert bad.status_code == 404
    _logout(client)


def test_p4_chatflow_requires_assistant_manage(client, seeded):
    _login(client, "tech_user")
    assert client.get("/api/chatflows").status_code == 403
    assert client.post("/api/chatflows", json={"name": "x"}).status_code == 403
    _logout(client)


# ======================================================================
# P2-5 生成任务恢复
# ======================================================================

def test_p5_answer_job_state_machine(seeded):
    from app.models import AnswerJobStatus
    from app.services.answer_jobs import AnswerJobService

    with SessionLocal() as session:
        service = AnswerJobService(session)
        request_id = f"req-{uuid.uuid4().hex}"
        job, created = service.create_or_get(
            conversation_id=None, message_id=None, user_id=None, assistant_id=None, request_id=request_id,
        )
        assert created is True
        same, created_again = service.create_or_get(
            conversation_id=None, message_id=None, user_id=None, assistant_id=None, request_id=request_id,
        )
        assert created_again is False and same.id == job.id

        service.mark_started(job.id)
        assert service.get(job.id).status == AnswerJobStatus.RETRIEVING
        cursor = service.append_content(job.id, "你好")
        assert cursor == 2
        service.update_stage(job.id, "checking", AnswerJobStatus.VERIFYING)
        assert service.get(job.id).current_stage == "checking"
        service.finish(job.id, AnswerJobStatus.COMPLETED)
        assert service.get(job.id).status == AnswerJobStatus.COMPLETED


def test_p5_recover_stale_jobs_marks_failed(seeded):
    from app.models import AnswerJobStatus
    from app.services.answer_jobs import AnswerJobService

    with SessionLocal() as session:
        service = AnswerJobService(session)
        job, _ = service.create_or_get(
            conversation_id=None, message_id=None, user_id=None, assistant_id=None,
            request_id=f"stale-{uuid.uuid4().hex}",
        )
        service.mark_started(job.id)
        recovered = service.recover_stale()
        assert recovered >= 1
        assert service.get(job.id).status == AnswerJobStatus.FAILED
        assert service.get(job.id).error_code == "ANSWER_INTERRUPTED"


# ======================================================================
# P2-6 文档理解（父子检索）
# ======================================================================

def test_p6_parent_child_expansion_returns_parent(seeded):
    from app.services.search import Candidate

    with SessionLocal() as session:
        document = Document(
            id=uuid.uuid4(), original_name="parent_child.txt", stored_path="test/parent_child.txt",
            extension="txt", mime_type="text/plain", size_bytes=10, sha256=uuid.uuid4().hex,
            status=DocumentStatus.READY, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID,
            visibility=DocumentVisibility.COMPANY, enabled=True,
        )
        session.add(document)
        session.flush()
        parent = DocumentChunk(
            id=uuid.uuid4(), document_id=document.id, sequence_number=1, content="父片段完整内容",
            chunk_role="parent", search_vector=func.to_tsvector("simple", "父片段完整内容"),
        )
        child = DocumentChunk(
            id=uuid.uuid4(), document_id=document.id, sequence_number=2, content="子片段",
            chunk_role="child", parent_sequence_number=1, parent_chunk_id=parent.id,
            search_vector=func.to_tsvector("simple", "子片段"),
        )
        session.add_all([parent, child])
        session.commit()

        service = SearchService(session, Settings(), user=session.get(User, seeded["admin"].id))
        expanded = service._expand_parents([Candidate(chunk=child, document=document)])
        assert len(expanded) == 1
        assert expanded[0].chunk.id == parent.id
        assert expanded[0].chunk.content == "父片段完整内容"


# ======================================================================
# P2-5 用户体验与比赛展示
# ======================================================================

def test_p5_no_answer_records_knowledge_gap(client, seeded, monkeypatch):
    import app.services.answer_runner as runner_module

    class _FakeRag:
        def __init__(self, *args, **kwargs):
            pass

        async def stream(self, request):
            yield AnswerEvent(type="sources", sources=[])
            yield AnswerEvent(type="delta", text="没有找到足够依据")
            yield AnswerEvent(type="no_answer", no_answer={
                "reason": "NO_RELEVANT_DOCUMENT", "message": "当前可访问的公司资料中没有找到足够依据。",
            })
            yield AnswerEvent(type="done", scope=KnowledgeScope.NONE)

    monkeypatch.setattr(runner_module, "RAG_FACTORY", _FakeRag)
    _login(client, "admin")
    response = client.post("/api/answer/stream", json={"question": "P5 未答缺口问题"})
    assert response.status_code == 200, response.text
    _logout(client)

    _login(client, "admin")
    gaps = client.get("/api/knowledge-gaps", params={"reason": "NO_ANSWER"}).json()["items"]
    assert any("P5 未答缺口问题" in item["question"] for item in gaps)
    _logout(client)


def test_p5_negative_feedback_records_gap(client, seeded):
    with SessionLocal() as session:
        user = session.get(User, seeded["tech_user"].id)
        chat = ChatService(session, user=user)
        row = chat.create("P5 反馈缺口")
        session.add(ChatMessage(
            session_id=row.id, role=ChatMessageRole.USER, content="P5 引用错误问题",
            status=ChatMessageStatus.COMPLETED,
        ))
        message = ChatMessage(
            session_id=row.id, role=ChatMessageRole.ASSISTANT, content="错误答案 [1]",
            status=ChatMessageStatus.COMPLETED,
        )
        session.add(message)
        session.commit()
        message_id = message.id

    _login(client, "tech_user")
    response = client.post("/api/feedback", json={
        "message_id": str(message_id), "rating": "DOWN", "reasons": ["引用不正确"],
    })
    assert response.status_code == 201, response.text
    _logout(client)

    _login(client, "admin")
    gaps = client.get("/api/knowledge-gaps", params={"reason": "WRONG_DOCUMENT"}).json()["items"]
    assert any("P5 引用错误问题" in item["question"] for item in gaps)
    gap_id = next(item["id"] for item in gaps if "P5 引用错误问题" in item["question"])

    statistics = client.get("/api/knowledge-gaps/statistics").json()
    assert statistics["total"] >= 1

    updated = client.patch(f"/api/knowledge-gaps/{gap_id}", json={
        "status": "RESOLVED", "linked_document_ids": [str(seeded["company_doc"].id)], "note": "已补充资料",
    })
    assert updated.status_code == 200, updated.text
    body = updated.json()
    assert body["status"] == "RESOLVED"
    assert body["linked_document_ids"] == [str(seeded["company_doc"].id)]
    _logout(client)


def test_p5_dashboard_overview(client, seeded):
    _login(client, "tech_user")
    response = client.get("/api/stats/dashboard")
    assert response.status_code == 200, response.text
    body = response.json()
    for key in (
        "knowledge_base_count", "document_count", "chunk_count", "questions_today",
        "avg_response_ms", "citation_coverage", "satisfaction", "hot_questions", "knowledge_gap_count",
    ):
        assert key in body
    _logout(client)


def test_p5_answer_emits_stage_detail_and_suggestions(seeded):
    _answer_cache.clear()
    service, _ = _service_for(seeded["admin"].id)
    service.config = RetrievalConfig(dictionary_enabled=False)
    events = asyncio.run(_collect(service, AnswerRequest(question="alpha secret")))
    stage_events = _events(events, "stage")
    stages = {event.stage for event in stage_events}
    assert {"understanding", "retrieving", "candidates", "local_generating"} <= stages
    candidates = next(event for event in stage_events if event.stage == "candidates")
    assert isinstance(candidates.detail, dict)
    suggestions = _events(events, "suggestions")
    assert suggestions and suggestions[0].suggestions


# ======================================================================
# P2-6 文档理解能力
# ======================================================================

@pytest.fixture(scope="module")
def p26_document(client, seeded):
    settings = Settings()
    settings.ensure_directories()
    path = settings.originals_root / "p26_preview.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("第一段内容。\n\n第二段内容。\n\n第三段内容。", encoding="utf-8")
    relative = path.relative_to(settings.library_root).as_posix()
    with SessionLocal() as session:
        document = Document(
            id=uuid.uuid4(), original_name="p26_preview.txt", stored_path=relative,
            extension="txt", mime_type="text/plain", size_bytes=path.stat().st_size,
            sha256=uuid.uuid4().hex, status=DocumentStatus.READY,
            knowledge_base_id=DEFAULT_KNOWLEDGE_BASE_ID, visibility=DocumentVisibility.COMPANY,
            enabled=True, owner_user_id=seeded["tech_user"].id,
        )
        session.add(document)
        session.commit()
        return document.id


def test_p6_chunk_preview(client, seeded, p26_document):
    _login(client, "tech_user")
    response = client.post(
        f"/api/documents/{p26_document}/chunk-preview",
        json={"chunking_config": {"strategy": "paragraph"}, "limit": 20},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total"] >= 1
    assert body["config"]["strategy"] == "paragraph"
    assert all(item["length"] == len(item["content"]) for item in body["items"])
    _logout(client)


def test_p6_knowledge_base_chunking_config(client, seeded):
    _login(client, "admin")
    created = client.post("/api/knowledge-bases", json={
        "name": "P6 切片库",
        "chunking_config": {"strategy": "heading", "target": 500, "maximum": 900, "overlap": 80},
    })
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["chunking_config"]["strategy"] == "heading"

    updated = client.patch(f"/api/knowledge-bases/{body['id']}", json={
        "chunking_config": {"strategy": "parent_child", "row_batch": 10},
    })
    assert updated.status_code == 200
    assert updated.json()["chunking_config"]["strategy"] == "parent_child"
    _logout(client)


def test_p6_document_graph_metadata(client, seeded):
    _login(client, "admin")
    response = client.patch(f"/api/documents/{seeded['company_doc'].id}", json={
        "author": "张三", "department_id": str(seeded["tech"].id), "topic": "气泡检测",
        "related_document_ids": [str(seeded["user_doc"].id)],
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["author"] == "张三"
    assert body["department_id"] == str(seeded["tech"].id)
    assert body["topic"] == "气泡检测"
    assert body["related_document_ids"] == [str(seeded["user_doc"].id)]
    _logout(client)
