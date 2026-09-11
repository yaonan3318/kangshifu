"""评测集服务：管理标准问题、版本化检索配置，批量运行并对比新旧版本。

P2-0 的目标是先建立可复现的质量基线：
- 评测集承载标准问题、答案关键点、正确/必须引用/禁止召回文档；
- 检索配置版本保存影响召回与精排的参数快照；
- 每次运行记录配置快照与逐用例明细，支持同一评测集上的新旧版本对比。
"""

import csv
import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import Settings
from app.errors import AppError
from app.models import (
    Document, DocumentChunk, EvaluationSet, KnowledgeBase, RetrievalConfigVersion,
    RetrievalTestCase, RetrievalTestRun,
)
from app.schemas.answer import AnswerRequest
from app.schemas.retrieval_lab import (
    ConfigVersionCreate, ConfigVersionUpdate, EvaluationSetCreate, EvaluationSetUpdate,
    RunCreate, TestCaseCreate, TestCaseUpdate,
)
from app.schemas.search import SearchRequest, SearchResult
from app.services.rag import RagService
from app.services.feedback_triggers import trigger_reverify
from app.services.retrieval_config import RetrievalConfig, config_diff, validate_feedback_config
from app.services.search import SearchService

METRIC_KEYS = (
    "document_recall", "chunk_recall", "hit_rate_at_1", "hit_rate_at_3", "hit_rate_at_5",
    "keypoint_coverage", "citation_accuracy", "no_answer_accuracy",
)

# 耗时指标单独保存；越低越好，用于配置/运行对比时展示响应时间变化。
LATENCY_KEYS = ("retrieval_ms", "first_token_latency_ms", "answer_latency_ms")

METRIC_LABELS: dict[str, str] = {
    "document_recall": "正确文档召回率",
    "chunk_recall": "正确片段召回率",
    "hit_rate_at_1": "Top1 命中率",
    "hit_rate_at_3": "Top3 命中率",
    "hit_rate_at_5": "Top5 命中率",
    "keypoint_coverage": "答案关键点覆盖率",
    "citation_accuracy": "引用正确率",
    "no_answer_accuracy": "无答案判断准确率",
    "retrieval_ms": "检索耗时(ms)",
    "first_token_latency_ms": "首字耗时(ms)",
    "answer_latency_ms": "完整回答耗时(ms)",
    "forbidden_violation_count": "禁止召回违规数",
}

# 这些指标数值越低越好；其余越高越好。
_LOWER_IS_BETTER = {"retrieval_ms", "first_token_latency_ms", "answer_latency_ms", "forbidden_violation_count"}

# 导入表头别名（中英文均可）。
HEADER_ALIASES: dict[str, set[str]] = {
    "name": {"name", "用例名称", "名称", "问题名称"},
    "question": {"question", "问题", "标准问题"},
    "knowledge_base": {"knowledge_base", "知识库", "知识库名称"},
    "expected_documents": {
        "expected_documents", "expected_document_ids", "正确文档", "预期文档", "应召回文档",
    },
    "expected_chunks": {
        "expected_chunks", "expected_chunk_ids", "期望片段", "预期片段", "正确片段", "应召回片段",
    },
    "must_cite_documents": {"must_cite_documents", "必须引用文档", "必须引用"},
    "forbidden_documents": {
        "forbidden_documents", "excluded_documents", "excluded_document_ids",
        "禁止召回文档", "不应召回", "禁止文档",
    },
    "answer_keypoints": {
        "answer_keypoints", "expected_answer_points", "答案关键点", "关键点", "期望答案关键点",
    },
    "expected_no_answer": {"expected_no_answer", "应无答案", "无答案"},
}
_TRUE_VALUES = {"1", "true", "yes", "y", "t", "是", "无答案", "应无答案"}


def _ratio(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 1.0
    return round(numerator / denominator, 4)


def _mean(values: list[float]) -> float:
    return round(sum(values) / len(values), 4) if values else 0.0


class RetrievalEvaluationService:
    def __init__(self, session: Session, settings: Settings, user=None):
        self.session = session
        self.settings = settings
        self.user = user
        self.search = SearchService(session, settings, user=user)
        # 允许测试注入替身，避免批量评测时访问真实模型。
        self.rag_factory = RagService

    # ------------------------------------------------------------ 评测集

    def list_evaluation_sets(self) -> list[EvaluationSet]:
        return list(self.session.scalars(select(EvaluationSet).order_by(EvaluationSet.created_at)))

    def get_evaluation_set(self, set_id: uuid.UUID) -> EvaluationSet:
        value = self.session.get(EvaluationSet, set_id)
        if value is None:
            raise AppError("EVALUATION_SET_NOT_FOUND", "评测集不存在", 404)
        return value

    def evaluation_set_payload(self, value: EvaluationSet) -> dict:
        count = self.session.scalar(
            select(func.count()).select_from(RetrievalTestCase)
            .where(RetrievalTestCase.evaluation_set_id == value.id)
        ) or 0
        return {
            "id": value.id, "name": value.name, "description": value.description,
            "knowledge_base_id": value.knowledge_base_id, "enabled": value.enabled,
            "case_count": count, "created_at": value.created_at, "updated_at": value.updated_at,
        }

    def create_evaluation_set(self, body: EvaluationSetCreate) -> EvaluationSet:
        if body.knowledge_base_id is not None:
            self._require_knowledge_base(body.knowledge_base_id)
        value = EvaluationSet(
            name=body.name.strip(), description=body.description,
            knowledge_base_id=body.knowledge_base_id, enabled=body.enabled,
        )
        self.session.add(value)
        self.session.commit()
        self.session.refresh(value)
        return value

    def update_evaluation_set(self, set_id: uuid.UUID, body: EvaluationSetUpdate) -> EvaluationSet:
        value = self.get_evaluation_set(set_id)
        supplied = body.model_fields_set
        if "knowledge_base_id" in supplied and body.knowledge_base_id is not None:
            self._require_knowledge_base(body.knowledge_base_id)
        for field in ("name", "description", "knowledge_base_id", "enabled"):
            if field in supplied:
                setattr(value, field, getattr(body, field))
        self.session.commit()
        self.session.refresh(value)
        return value

    def delete_evaluation_set(self, set_id: uuid.UUID) -> None:
        value = self.get_evaluation_set(set_id)
        self.session.delete(value)
        self.session.commit()

    # ------------------------------------------------------------ 标准问题

    def list_cases(self, set_id: uuid.UUID | None = None) -> list[RetrievalTestCase]:
        statement = select(RetrievalTestCase)
        if set_id is not None:
            statement = statement.where(RetrievalTestCase.evaluation_set_id == set_id)
        return list(self.session.scalars(statement.order_by(RetrievalTestCase.created_at.desc())))

    def get_case(self, case_id: uuid.UUID) -> RetrievalTestCase:
        value = self.session.get(RetrievalTestCase, case_id)
        if not value:
            raise AppError("RETRIEVAL_CASE_NOT_FOUND", "标准问题不存在", 404)
        return value

    def create_case(self, body: TestCaseCreate, set_id: uuid.UUID | None = None) -> RetrievalTestCase:
        if set_id is not None:
            self.get_evaluation_set(set_id)
        self._validate_documents([
            *body.expected_document_ids, *body.must_cite_document_ids, *body.forbidden_document_ids,
        ])
        self._validate_chunks(body.expected_chunk_ids)
        value = RetrievalTestCase(
            evaluation_set_id=set_id,
            name=body.name.strip(), question=body.question.strip(), knowledge_base_id=body.knowledge_base_id,
            expected_document_ids=[str(item) for item in body.expected_document_ids],
            expected_chunk_ids=[str(item) for item in body.expected_chunk_ids],
            must_cite_document_ids=[str(item) for item in body.must_cite_document_ids],
            forbidden_document_ids=[str(item) for item in body.forbidden_document_ids],
            expected_keywords=[item.strip() for item in body.expected_keywords if item.strip()],
            expected_answer_keypoints=[item.strip() for item in body.expected_answer_keypoints if item.strip()],
            expected_no_answer=body.expected_no_answer, enabled=body.enabled,
        )
        self.session.add(value)
        self.session.commit()
        self.session.refresh(value)
        return value

    def update_case(self, case_id: uuid.UUID, body: TestCaseUpdate) -> RetrievalTestCase:
        value = self.get_case(case_id)
        supplied = body.model_fields_set
        for field in ("expected_document_ids", "must_cite_document_ids", "forbidden_document_ids"):
            if field in supplied and getattr(body, field) is not None:
                self._validate_documents(getattr(body, field))
                setattr(value, field, [str(item) for item in getattr(body, field)])
        if "expected_chunk_ids" in supplied and body.expected_chunk_ids is not None:
            self._validate_chunks(body.expected_chunk_ids)
            value.expected_chunk_ids = [str(item) for item in body.expected_chunk_ids]
        for field in ("name", "question"):
            if field in supplied and getattr(body, field) is None:
                raise AppError("INVALID_RETRIEVAL_CASE", f"{field} 不能为空", 400)
        for field in ("name", "question", "knowledge_base_id", "expected_no_answer", "enabled"):
            if field in supplied:
                setattr(value, field, getattr(body, field))
        for field in ("expected_keywords", "expected_answer_keypoints"):
            if field in supplied and getattr(body, field) is not None:
                setattr(value, field, [item.strip() for item in getattr(body, field) if item.strip()])
        self.session.commit()
        self.session.refresh(value)
        return value

    def delete_case(self, case_id: uuid.UUID) -> None:
        self.session.delete(self.get_case(case_id))
        self.session.commit()

    # ------------------------------------------------------------ 批量导入

    def import_cases(self, set_id: uuid.UUID, filename: str, content: bytes) -> dict:
        self.get_evaluation_set(set_id)
        rows = self._parse_import(filename, content)
        created = skipped = 0
        errors: list[str] = []
        for index, raw in enumerate(rows, start=2):
            try:
                self._import_row(set_id, raw)
                created += 1
            except AppError as exc:
                skipped += 1
                errors.append(f"第 {index} 行：{exc.message}")
            except Exception as exc:  # pragma: no cover - 防御解析异常
                skipped += 1
                errors.append(f"第 {index} 行：{exc}")
        return {"created": created, "skipped": skipped, "errors": errors[:50]}

    def _parse_import(self, filename: str, content: bytes) -> list[dict[str, str]]:
        lowered = (filename or "").lower()
        if lowered.endswith(".csv"):
            text = content.decode("utf-8-sig", errors="replace")
            return [dict(row) for row in csv.DictReader(io.StringIO(text))]
        if lowered.endswith((".xlsx", ".xlsm")):
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                return []
            columns = [str(cell).strip() if cell is not None else "" for cell in header]
            parsed = []
            for row in rows:
                parsed.append({columns[i]: ("" if row[i] is None else str(row[i])) for i in range(min(len(columns), len(row)))})
            return parsed
        raise AppError("UNSUPPORTED_IMPORT_FORMAT", "只支持 CSV 或 Excel(.xlsx) 文件", 415)

    def _import_row(self, set_id: uuid.UUID, raw: dict[str, str]) -> RetrievalTestCase:
        data = self._normalize_headers(raw)
        question = (data.get("question") or "").strip()
        if not question:
            raise AppError("IMPORT_MISSING_QUESTION", "缺少问题字段", 400)
        knowledge_base_id = self._resolve_knowledge_base(data.get("knowledge_base"))
        expected = self._resolve_documents(data.get("expected_documents"), knowledge_base_id)
        expected_chunks = self._resolve_chunks(data.get("expected_chunks"))
        must_cite = self._resolve_documents(data.get("must_cite_documents"), knowledge_base_id)
        forbidden = self._resolve_documents(data.get("forbidden_documents"), knowledge_base_id)
        keypoints = self._split_tokens(data.get("answer_keypoints"))
        no_answer = (data.get("expected_no_answer") or "").strip().lower() in _TRUE_VALUES
        name = (data.get("name") or "").strip() or question[:40]
        return self.create_case(TestCaseCreate(
            name=name, question=question, knowledge_base_id=knowledge_base_id,
            expected_document_ids=expected, expected_chunk_ids=expected_chunks,
            must_cite_document_ids=must_cite,
            forbidden_document_ids=forbidden, expected_answer_keypoints=keypoints,
            expected_no_answer=no_answer,
        ), set_id=set_id)

    @staticmethod
    def _normalize_headers(raw: dict[str, str]) -> dict[str, str]:
        lookup = {key.strip().lower(): value for key, value in raw.items() if key}
        result: dict[str, str] = {}
        for canonical, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                if alias.lower() in lookup:
                    result[canonical] = lookup[alias.lower()]
                    break
        return result

    @staticmethod
    def _split_tokens(value: str | None) -> list[str]:
        if not value:
            return []
        normalized = value.replace("|", ";").replace("，", ",").replace("；", ";").replace(",", ";")
        return [item.strip() for item in normalized.split(";") if item.strip()]

    def _resolve_knowledge_base(self, value: str | None) -> uuid.UUID | None:
        token = (value or "").strip()
        if not token:
            return None
        try:
            return uuid.UUID(token)
        except ValueError:
            pass
        row = self.session.scalar(select(KnowledgeBase.id).where(KnowledgeBase.name == token))
        if row is None:
            raise AppError("IMPORT_KNOWLEDGE_BASE_NOT_FOUND", f"知识库不存在：{token}", 400)
        return row

    def _resolve_documents(self, value: str | None, knowledge_base_id: uuid.UUID | None) -> list[uuid.UUID]:
        result: list[uuid.UUID] = []
        for token in self._split_tokens(value):
            try:
                result.append(uuid.UUID(token))
                continue
            except ValueError:
                pass
            clauses = [Document.original_name == token, Document.deleted_at.is_(None)]
            if knowledge_base_id is not None:
                clauses.append(Document.knowledge_base_id == knowledge_base_id)
            document_id = self.session.scalar(select(Document.id).where(*clauses).limit(1))
            if document_id is None:
                raise AppError("IMPORT_DOCUMENT_NOT_FOUND", f"文档不存在：{token}", 400)
            result.append(document_id)
        return list(dict.fromkeys(result))

    def _resolve_chunks(self, value: str | None) -> list[uuid.UUID]:
        """解析期望片段；只接受 chunk_id（UUID），并校验片段确实存在。"""
        result: list[uuid.UUID] = []
        for token in self._split_tokens(value):
            try:
                result.append(uuid.UUID(token))
            except ValueError as exc:
                raise AppError("IMPORT_CHUNK_NOT_FOUND", f"片段 ID 不是有效 UUID：{token}", 400) from exc
        result = list(dict.fromkeys(result))
        if result:
            self._validate_chunks(result)
        return result

    # ------------------------------------------------------------ 配置版本

    def list_config_versions(self) -> list[RetrievalConfigVersion]:
        return list(self.session.scalars(
            select(RetrievalConfigVersion).order_by(RetrievalConfigVersion.created_at.desc())
        ))

    def get_config_version(self, config_id: uuid.UUID) -> RetrievalConfigVersion:
        value = self.session.get(RetrievalConfigVersion, config_id)
        if value is None:
            raise AppError("RETRIEVAL_CONFIG_NOT_FOUND", "检索配置版本不存在", 404)
        return value

    def _default_config_version(self) -> RetrievalConfigVersion | None:
        return self.session.scalar(
            select(RetrievalConfigVersion).where(RetrievalConfigVersion.is_default.is_(True)).limit(1)
        )

    def create_config_version(self, body: ConfigVersionCreate) -> RetrievalConfigVersion:
        normalized = RetrievalConfig.from_dict(body.config, base=RetrievalConfig.from_settings(self.settings))
        validate_feedback_config(normalized)
        value = RetrievalConfigVersion(
            name=body.name.strip(), description=body.description,
            config=normalized.to_dict(), is_default=False,
        )
        self.session.add(value)
        self.session.flush()
        if body.is_default:
            self._set_default(value)
        self.session.commit()
        self.session.refresh(value)
        trigger_reverify(
            self.session, config_version_id=value.id, reason="retrieval_config_created",
        )
        return value

    def update_config_version(self, config_id: uuid.UUID, body: ConfigVersionUpdate) -> RetrievalConfigVersion:
        value = self.get_config_version(config_id)
        supplied = body.model_fields_set
        if "name" in supplied and body.name is not None:
            value.name = body.name.strip()
        if "description" in supplied:
            value.description = body.description
        if "config" in supplied and body.config is not None:
            normalized = RetrievalConfig.from_dict(
                body.config, base=RetrievalConfig.from_settings(self.settings),
            )
            validate_feedback_config(normalized)
            value.config = normalized.to_dict()
        if body.is_default:
            self._set_default(value)
        self.session.commit()
        self.session.refresh(value)
        trigger_reverify(
            self.session, config_version_id=value.id, reason="retrieval_config_updated",
        )
        return value

    def delete_config_version(self, config_id: uuid.UUID) -> None:
        value = self.get_config_version(config_id)
        self.session.delete(value)
        self.session.commit()

    def _set_default(self, value: RetrievalConfigVersion) -> None:
        for other in self.session.scalars(
            select(RetrievalConfigVersion).where(RetrievalConfigVersion.is_default.is_(True))
        ):
            other.is_default = False
        value.is_default = True

    def config_diff(self, left_id: uuid.UUID, right_id: uuid.UUID) -> list[dict]:
        left = RetrievalConfig.from_dict(self.get_config_version(left_id).config)
        right = RetrievalConfig.from_dict(self.get_config_version(right_id).config)
        return config_diff(left, right)

    # ------------------------------------------------------------ 运行评测

    async def run_evaluation(self, body: RunCreate) -> RetrievalTestRun:
        evaluation_set = self.get_evaluation_set(body.evaluation_set_id)
        config_version = (
            self.get_config_version(body.config_version_id) if body.config_version_id
            else self._default_config_version()
        )
        config = RetrievalConfig.from_dict(
            config_version.config if config_version else {}, base=RetrievalConfig.from_settings(self.settings),
        )
        cases = [case for case in self.list_cases(evaluation_set.id) if case.enabled]
        if not cases:
            raise AppError("NO_RETRIEVAL_CASES", "评测集没有已启用的标准问题", 409)
        outcomes = [
            await self._run_case(case, evaluation_set, config, body.limit, body.include_answers)
            for case in cases
        ]
        metrics = self._aggregate(outcomes, body.include_answers)
        run = RetrievalTestRun(
            evaluation_set_id=evaluation_set.id,
            config_version_id=config_version.id if config_version else None,
            settings_snapshot=self._settings_snapshot(),
            config_snapshot=config.to_dict(),
            results=outcomes, metrics=metrics,
        )
        self.session.add(run)
        self.session.commit()
        self.session.refresh(run)
        return run

    async def _run_case(
        self, case: RetrievalTestCase, evaluation_set: EvaluationSet,
        config: RetrievalConfig, limit: int, include_answers: bool,
    ) -> dict:
        knowledge_base_id = case.knowledge_base_id or evaluation_set.knowledge_base_id
        search = SearchService(self.session, self.settings, user=self.user, config=config)
        outcome = search.search_with_diagnostics(SearchRequest(
            query=case.question, knowledge_base_id=knowledge_base_id, limit=limit,
        ))
        result = self._score_case(case, outcome.items, outcome.diagnostics)
        if include_answers:
            answer = await self._generate_answer(case, knowledge_base_id, config)
            result.update(answer)
        return result

    @staticmethod
    def _score_case(case: RetrievalTestCase, returned: list[SearchResult], diagnostics=None) -> dict:
        """按“片段级”真实命中计算逐用例指标；检索阶段不产生首字/回答耗时。"""
        returned_doc_ids = list(dict.fromkeys(str(item.document_id) for item in returned))
        returned_chunk_ids = [str(item.chunk_id) for item in returned]
        expected_docs = {str(item) for item in case.expected_document_ids}
        expected_chunks = {str(item) for item in case.expected_chunk_ids}
        must_cite = {str(item) for item in case.must_cite_document_ids}
        forbidden = {str(item) for item in case.forbidden_document_ids}

        matched_chunks = expected_chunks & set(returned_chunk_ids)
        # 片段召回率 = 命中的期望片段数 / 期望片段总数；未配置期望片段时视为满分。
        chunk_recall = _ratio(len(matched_chunks), len(expected_chunks))
        # Top-K 命中率优先按片段计算；没有期望片段时退回文档级，避免指标失真。
        if expected_chunks:
            hit1 = RetrievalEvaluationService._hit_rate(expected_chunks, returned_chunk_ids, 1)
            hit3 = RetrievalEvaluationService._hit_rate(expected_chunks, returned_chunk_ids, 3)
            hit5 = RetrievalEvaluationService._hit_rate(expected_chunks, returned_chunk_ids, 5)
        else:
            hit1 = RetrievalEvaluationService._hit_rate(expected_docs, returned_doc_ids, 1)
            hit3 = RetrievalEvaluationService._hit_rate(expected_docs, returned_doc_ids, 3)
            hit5 = RetrievalEvaluationService._hit_rate(expected_docs, returned_doc_ids, 5)
        timings = diagnostics.timings_ms if diagnostics is not None else {}
        evidence = "\n".join(item.content.lower() for item in returned)
        return {
            "case_id": str(case.id), "name": case.name, "question": case.question,
            "returned_document_ids": returned_doc_ids,
            "returned_chunk_ids": returned_chunk_ids,
            "expected_chunk_ids": sorted(expected_chunks),
            "matched_chunk_ids": sorted(matched_chunks),
            "forbidden_hits": [item for item in returned_doc_ids if item in forbidden],
            "document_recall": _ratio(len(expected_docs & set(returned_doc_ids)), len(expected_docs)),
            "chunk_recall": chunk_recall,
            "hit_rate_at_1": hit1, "hit_rate_at_3": hit3, "hit_rate_at_5": hit5,
            "citation_accuracy": _ratio(len(must_cite & set(returned_doc_ids)), len(must_cite)),
            "no_answer_accuracy": 1.0 if case.expected_no_answer == (not returned) else 0.0,
            "keypoint_coverage": RetrievalEvaluationService._keypoint_coverage(case.expected_answer_keypoints, evidence),
            # 检索阶段还没有生成答案，首字/完整回答耗时记为 0，由回答阶段覆盖。
            "retrieval_ms": timings.get("total", 0.0),
            "first_token_latency_ms": 0.0,
            "answer_latency_ms": 0.0,
            "answer_mode": "evidence",
            "mode": diagnostics.mode if diagnostics is not None else None,
            "warning": diagnostics.warning if diagnostics is not None else None,
        }

    @staticmethod
    def _hit_rate(expected: set[str], returned: list[str], k: int) -> float:
        if not expected:
            return 1.0
        return 1.0 if expected & set(returned[:k]) else 0.0

    async def _generate_answer(
        self, case: RetrievalTestCase, knowledge_base_id: uuid.UUID | None, config: RetrievalConfig,
    ) -> dict:
        """可选：运行完整 RAG 回答，得到真实关键点覆盖与引用正确率。"""
        # 评测必须每次真实生成，关闭回答缓存，避免复用得分的耗时数据。
        eval_settings = self.settings.model_copy(update={"answer_cache_enabled": False})
        service = self.rag_factory(
            self.session, eval_settings, user=self.user, config=config, source_limit=config.final_limit,
        )
        text_parts: list[str] = []
        sources: list = []
        metrics: dict = {}
        scope = None
        async for event in service.stream(AnswerRequest(question=case.question, knowledge_base_id=knowledge_base_id)):
            if event.type == "replace":
                text_parts.clear()
            elif event.type == "delta" and event.text:
                text_parts.append(event.text)
            elif event.type == "sources" and event.sources is not None:
                sources = event.sources
            elif event.type == "metrics" and event.metrics is not None:
                metrics = event.metrics
            elif event.type == "done" and event.scope is not None:
                scope = event.scope.value
        answer_text = "".join(text_parts)
        cited = list(dict.fromkeys(str(item.document_id) for item in sources))
        must_cite = {str(item) for item in case.must_cite_document_ids}
        no_answer = not sources or scope == "NONE"
        return {
            "answer_mode": "answer",
            "answer_text": answer_text,
            "cited_document_ids": cited,
            "keypoint_coverage": self._keypoint_coverage(case.expected_answer_keypoints, answer_text.lower()),
            "citation_accuracy": _ratio(len(must_cite & set(cited)), len(must_cite)),
            "no_answer_accuracy": 1.0 if case.expected_no_answer == no_answer else 0.0,
            "first_token_latency_ms": metrics.get("llm_first_token_ms") or 0.0,
            "answer_latency_ms": metrics.get("total_ms") or 0.0,
        }

    @staticmethod
    def _keypoint_coverage(keypoints: list[str], text: str) -> float:
        if not keypoints:
            return 1.0
        hits = sum(1 for keypoint in keypoints if keypoint.strip().lower() in text)
        return _ratio(hits, len(keypoints))

    @staticmethod
    def _aggregate(outcomes: list[dict], include_answers: bool) -> dict:
        metrics = {"case_count": len(outcomes), "answer_mode": "answer" if include_answers else "evidence"}
        for key in METRIC_KEYS:
            metrics[key] = _mean([float(item.get(key, 0.0)) for item in outcomes])
        for key in ("retrieval_ms", "first_token_latency_ms", "answer_latency_ms"):
            metrics[key] = round(sum(float(item.get(key, 0.0)) for item in outcomes) / len(outcomes), 2)
        metrics["forbidden_violation_count"] = sum(1 for item in outcomes if item["forbidden_hits"])
        return metrics

    def list_runs(self, set_id: uuid.UUID | None = None) -> list[RetrievalTestRun]:
        statement = select(RetrievalTestRun)
        if set_id is not None:
            statement = statement.where(RetrievalTestRun.evaluation_set_id == set_id)
        return list(self.session.scalars(statement.order_by(RetrievalTestRun.created_at.desc()).limit(100)))

    def get_run(self, run_id: uuid.UUID) -> RetrievalTestRun:
        value = self.session.get(RetrievalTestRun, run_id)
        if not value:
            raise AppError("RETRIEVAL_RUN_NOT_FOUND", "评测记录不存在", 404)
        return value

    def compare_runs(self, left_id: uuid.UUID, right_id: uuid.UUID) -> dict:
        left = self.get_run(left_id)
        right = self.get_run(right_id)
        deltas = {
            key: round(float(right.metrics.get(key, 0.0)) - float(left.metrics.get(key, 0.0)), 4)
            for key in METRIC_KEYS
        }
        left_config = RetrievalConfig.from_dict(left.config_snapshot)
        right_config = RetrievalConfig.from_dict(right.config_snapshot)
        changes = self._case_changes(left.results, right.results)
        return {
            "left": left, "right": right, "metric_deltas": deltas,
            "metric_changes": self._metric_changes(left.metrics, right.metrics),
            "config_differences": config_diff(left_config, right_config),
            "case_changes": changes,
        }

    @staticmethod
    def _metric_changes(left_metrics: dict, right_metrics: dict) -> list[dict]:
        """逐指标展示旧值、新值、差值、提升或下降，以及响应时间变化。"""
        keys = [*METRIC_KEYS, "forbidden_violation_count", *LATENCY_KEYS]
        result: list[dict] = []
        for key in keys:
            left_value = float(left_metrics.get(key, 0.0))
            right_value = float(right_metrics.get(key, 0.0))
            delta = round(right_value - left_value, 4)
            if delta == 0:
                direction, improved = "same", None
            elif key in _LOWER_IS_BETTER:
                direction = "down" if delta < 0 else "up"
                improved = delta < 0
            else:
                direction = "up" if delta > 0 else "down"
                improved = delta > 0
            result.append({
                "key": key, "label": METRIC_LABELS.get(key, key),
                "left": round(left_value, 4), "right": round(right_value, 4),
                "delta": delta, "direction": direction, "improved": improved,
                "unit": "ms" if key in LATENCY_KEYS else "ratio",
                "lower_is_better": key in _LOWER_IS_BETTER,
            })
        return result

    @staticmethod
    def _case_changes(left_results: list[dict], right_results: list[dict]) -> list[dict]:
        left_cases = {item.get("case_id"): item for item in left_results}
        right_cases = {item.get("case_id"): item for item in right_results}
        changes = []
        for case_id in sorted(set(left_cases) | set(right_cases)):
            left_case = left_cases.get(case_id, {})
            right_case = right_cases.get(case_id, {})
            changes.append({
                "case_id": case_id,
                "name": right_case.get("name") or left_case.get("name"),
                "left_document_recall": left_case.get("document_recall"),
                "right_document_recall": right_case.get("document_recall"),
                "document_recall_delta": round(
                    float(right_case.get("document_recall", 0.0)) - float(left_case.get("document_recall", 0.0)), 4,
                ),
                "left_chunk_recall": left_case.get("chunk_recall"),
                "right_chunk_recall": right_case.get("chunk_recall"),
                "chunk_recall_delta": round(
                    float(right_case.get("chunk_recall", 0.0)) - float(left_case.get("chunk_recall", 0.0)), 4,
                ),
                "left_no_answer_accuracy": left_case.get("no_answer_accuracy"),
                "right_no_answer_accuracy": right_case.get("no_answer_accuracy"),
                "left_answer_latency_ms": left_case.get("answer_latency_ms"),
                "right_answer_latency_ms": right_case.get("answer_latency_ms"),
            })
        return changes

    # ------------------------------------------------------------ 辅助

    def _validate_documents(self, document_ids) -> None:
        for document_id in document_ids:
            value = self.session.scalar(select(Document.id).where(Document.id == document_id, Document.deleted_at.is_(None)))
            if not value:
                raise AppError("EXPECTED_DOCUMENT_UNAVAILABLE", f"预期文档不存在或已删除：{document_id}", 400)

    def _validate_chunks(self, chunk_ids) -> None:
        for chunk_id in chunk_ids:
            value = self.session.scalar(select(DocumentChunk.id).where(DocumentChunk.id == chunk_id))
            if not value:
                raise AppError("EXPECTED_CHUNK_UNAVAILABLE", f"预期片段不存在：{chunk_id}", 400)

    def _require_knowledge_base(self, knowledge_base_id: uuid.UUID) -> None:
        if self.session.scalar(select(KnowledgeBase.id).where(KnowledgeBase.id == knowledge_base_id)) is None:
            raise AppError("KNOWLEDGE_BASE_NOT_FOUND", "知识库不存在", 404)

    def _settings_snapshot(self) -> dict:
        return {
            "embedding_model": self.settings.embedding_model,
            "candidate_limit": self.settings.search_candidate_limit,
            "rrf_k": self.settings.search_rrf_k,
            "vector_min_similarity": self.settings.search_vector_min_similarity,
            "min_evidence_score": self.settings.search_min_evidence_score,
            "rerank_enabled": self.settings.rerank_enabled,
            "rerank_model": self.settings.rerank_model,
        }
